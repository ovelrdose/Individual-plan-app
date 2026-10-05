"""Инструкторы и их постоянный распорядок из шахматки заказчика (``instructions_schedule.xlsx``)
— чистый Python без Django.

Из файла берутся только фамилии инструкторов (колонки, «A/ B» — пара 2/2) и то, что они
делают в слоте почти каждый день: «Метод. работа», БОС, ведение группы. Записи пациентов не
читаются и никуда не попадают (CLAUDE.md: персональные данные пациентов в систему не
переносятся).

Состав колонок за год менялся, поэтому берутся последние листы с тем же составом, что и на
последнем листе. Обязанность — постоянная, если в этом слоте она стоит не меньше чем в
половине этих дней. Группу, которую ведут попеременно, получает тот, кто ведёт её чаще.
"""

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import time
from pathlib import Path

from openpyxl import load_workbook

# Колонки тренажёров справа — не инструкторы.
NOT_INSTRUCTORS = {"st", "st-150", "имитрон", "имтрн", "pablo", "pb", "время"}
PATIENT = re.compile(r"^\s*\d+\s*[а-яa-z]?\s*(п\b|омр)|омр", re.IGNORECASE)
METHOD = "метод. работа"
BOS = "бос"


@dataclass(frozen=True)
class Duty:
    slot: time
    kind: str  # METHOD_WORK, BOS, GROUP_LEAD
    group: str = ""  # для GROUP_LEAD — название группы как в справочнике


@dataclass
class Column:
    """Колонка шахматки: одиночка или пара 2/2 и её постоянный распорядок."""

    members: list[str]
    duties: list[Duty] = field(default_factory=list)


@dataclass
class BoardStaff:
    columns: list[Column]
    sheets: int  # сколько листов с этим составом разобрано
    first_sheet: str
    last_sheet: str


def read_board_staff(path: Path, groups: dict[str, str]) -> BoardStaff:
    """``groups`` — нормализованное название группы → название в справочнике
    («i-can нога» → «I can нога»)."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = _header(rows[0])
        if header:
            sheets.append((sheet.title, header, rows[1:]))
    workbook.close()
    if not sheets:
        raise ValueError("В файле нет листов шахматки.")
    current = sheets[-1][1]
    recent = []
    for item in reversed(sheets):
        if item[1] != current:
            break
        recent.append(item)
    recent.reverse()

    counts: dict[tuple[int, time], Counter] = defaultdict(Counter)
    for _title, _header_row, rows in recent:
        for row in rows:
            slot = _slot(row[0] if row else None)
            if slot is None:
                continue
            for index in range(len(current)):
                value = row[index + 1] if index + 1 < len(row) else None
                label = _duty_label(value, groups)
                if label is not None:
                    counts[(index, slot)][label] += 1

    columns = [Column(_members(name)) for name in current]
    days = len(recent)
    leaders: dict[tuple[str, time], tuple[int, int]] = {}
    for (index, slot), counter in counts.items():
        for label, count in counter.items():
            if label in (METHOD, BOS):
                if 2 * count >= days:
                    kind = "METHOD_WORK" if label == METHOD else "BOS"
                    columns[index].duties.append(Duty(slot, kind))
            else:
                best = leaders.get((label, slot))
                if best is None or count > best[1]:
                    leaders[(label, slot)] = (index, count)
    for (group, slot), (index, _count) in leaders.items():
        total = sum(counts[(i, slot)][group] for i in range(len(current)))
        if 2 * total >= days:
            columns[index].duties.append(Duty(slot, "GROUP_LEAD", group))
    for column in columns:
        column.duties.sort(key=lambda d: (d.slot, d.kind))
    return BoardStaff(columns, days, recent[0][0], recent[-1][0])


def _header(row: tuple) -> tuple[str, ...]:
    names = []
    for value in row[1:]:
        text = " ".join(str(value).split()) if value is not None else ""
        if not text:
            continue
        if text.lower() in NOT_INSTRUCTORS or re.fullmatch(r"[\d.,]+", text):
            break
        names.append(_normalize_pair(text))
    return tuple(names)


def _normalize_pair(text: str) -> str:
    return "/ ".join(part.strip() for part in text.split("/"))


def _members(name: str) -> list[str]:
    return [part.strip() for part in name.split("/") if part.strip()]


def _slot(value) -> time | None:
    """«9:10-9:40» → 9:10; вечерние строки записаны числом: 18 → 18:00, 18.4 → 18:40."""
    if value is None:
        return None
    if isinstance(value, time):
        return value
    text = str(value).strip()
    match = re.match(r"^(\d{1,2}):(\d{2})", text)
    if match:
        return time(int(match[1]), int(match[2]))
    match = re.fullmatch(r"(\d{1,2})(?:[.,](\d))?", text)
    if match and 7 <= int(match[1]) <= 21:
        return time(int(match[1]), int(match[2] or 0) * 10)
    return None


def _duty_label(value, groups: dict[str, str]) -> str | None:
    """Обязанность в ячейке или None (пусто, пациент, прочее)."""
    if value is None:
        return None
    text = " ".join(str(value).split()).lower()
    if not text or PATIENT.search(text):
        return None
    if text == METHOD:
        return METHOD
    if text == BOS:
        return BOS
    key = normalize_group(text)
    return groups.get(key)


def normalize_group(name: str) -> str:
    """«I can нога», «i-can нога», «Вест.гимнастика» → сравнимый ключ."""
    text = name.lower().replace("-", " ").replace(".", " ")
    return " ".join(text.split())
