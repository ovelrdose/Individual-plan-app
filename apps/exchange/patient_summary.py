"""Разбор сводки пациентов поста (.docx, «Сводка пациентов ОМР 4 пост 1 на 15.10.26») — без Django.

Сводку ведут постовые сёстры в Word: таблица «палата | ФИО | возраст | пол | ИБ | стол | диагноз |
бюджет | дата поступления | дата выписки | койко-дни | лечащий врач». Берём только то, что нужно
программе; диагноз, стол и бюджет не читаем (решение 70).

Палата пишется по-разному: «5а», «8п», « / 3» (продолжение объединённой ячейки), «П / И / Т»
(палата интенсивной терапии), «8п(омр2)», «Омр3» с палатой в скобках после ФИО — «Печенин А.Я
(6 п)». Строки без даты поступления — сопровождающие, пустые и повторы объединённых ячеек.
"""

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import BinaryIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError

INTENSIVE_CARE = "ПИТ"


class SummaryError(ValueError):
    """Файл не похож на сводку пациентов. Текст показывается пользователю."""


@dataclass(frozen=True)
class SummaryPatient:
    full_name: str  # «Астахов М.И.» — в сводке только инициалы
    room: str  # «5а», «ПИТ»
    department: str  # «ОМР 3» — пациент другого отделения, иначе пусто
    sex: str  # «М», «Ж» или пусто
    age: int | None
    history_number: str
    admitted: date
    discharge: date | None  # пусто — дата выписки ещё не известна
    doctor: str  # «Петлина К.И.»

    @property
    def surname(self) -> str:
        return self.full_name.split()[0]


@dataclass
class Summary:
    day: date | None  # «на 15.10.26» из заголовка
    patients: list[SummaryPatient] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)  # ФИО без даты поступления


# Колонка → слова в её заголовке (заголовки в сводках постов немного разные).
COLUMNS = {
    "name": ("фио",),
    "age": ("в",),
    "sex": ("м/ж",),
    "history_number": ("иб",),
    "admitted": ("поступ",),
    "discharge": ("выпис",),
    "doctor": ("врач", "л/в"),
}
DATE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{2}(?:\d{2})?)")
ROOM_IN_NAME = re.compile(r"\(\s*(\d+\s*[а-яё]?)\s*п?\s*\)", re.IGNORECASE)
DEPARTMENT = re.compile(r"омр\s*№?\s*(\d+)", re.IGNORECASE)


def read_summary(source: BinaryIO | Path | str) -> Summary:
    try:
        document = Document(source)
    except (PackageNotFoundError, BadZipFile, XMLSyntaxError, KeyError, ValueError) as error:
        raise SummaryError("Файл не читается как документ Word (.docx).") from error
    day = next((_date(m.group(0)) for p in document.paragraphs if (m := DATE.search(p.text))), None)
    for table in document.tables:
        rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        for index, row in enumerate(rows):
            columns = _columns(row)
            if columns is not None:
                summary = Summary(day)
                for cells in rows[index + 1 :]:
                    _add_row(summary, cells, columns)
                return summary
    raise SummaryError("В документе нет таблицы пациентов с колонкой «ФИО».")


def _clean(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split())


def _columns(header: list[str]) -> dict[str, int] | None:
    found: dict[str, int] = {}
    for index, text in enumerate(header):
        words = text.lower().replace(" / ", " ").replace("№", " ").split()
        for key, marks in COLUMNS.items():
            if key in found:
                continue
            if any(mark in words or (len(mark) > 2 and mark in text.lower()) for mark in marks):
                found[key] = index
                break
    return found if found.keys() == COLUMNS.keys() else None


def _add_row(summary: Summary, cells: list[str], columns: dict[str, int]) -> None:
    def cell(key: str) -> str:
        index = columns[key]
        return cells[index] if index < len(cells) else ""

    raw_name = cell("name")
    if not raw_name or raw_name.startswith("+"):
        return
    admitted = _date(cell("admitted"))
    if admitted is None:
        name = _person(raw_name)
        if name and name not in summary.skipped and not _known(summary, name):
            summary.skipped.append(name)
        return
    number = cell("history_number")
    if number and any(p.history_number == number for p in summary.patients):
        return  # повтор объединённой ячейки
    place = cells[0] if cells else ""
    age = cell("age")
    sex = cell("sex").upper()[:1]
    summary.patients.append(
        SummaryPatient(
            full_name=_person(raw_name),
            room=_room(place, raw_name),
            department=_department(place, raw_name),
            sex=sex if sex in ("М", "Ж") else "",
            age=int(age) if age.isdigit() else None,
            history_number=number,
            admitted=admitted,
            discharge=_date(cell("discharge")),
            doctor=_person(cell("doctor")),
        )
    )


def _known(summary: Summary, name: str) -> bool:
    return any(p.full_name == name for p in summary.patients)


def _person(text: str) -> str:
    """«Шушков В. В. ( 9 п)», «Плаунов С.А. +сопр», «Андрусенко А,А,» → «Фамилия И.О.»."""
    text = re.sub(r"\(.*?\)|\+.*", " ", text).replace(",", ".")
    parts = text.split()
    if not parts:
        return ""
    initials = re.findall(r"[А-ЯЁA-Z]", "".join(parts[1:]))
    return parts[0] + (" " + "".join(f"{letter}." for letter in initials) if initials else "")


def _room(place: str, name: str) -> str:
    compact = place.replace("/", "").replace(" ", "")
    if compact.upper() == "ПИТ":
        return INTENSIVE_CARE
    if in_name := ROOM_IN_NAME.search(name):
        compact = in_name.group(1).replace(" ", "")
    else:
        compact = DEPARTMENT.sub("", re.sub(r"\(.*?\)", "", compact))
    # «8п» — «палата 8»; «5а» — вариант палаты (глоссарий).
    return re.sub(r"п$", "", compact, flags=re.IGNORECASE)


def _department(place: str, name: str) -> str:
    found = DEPARTMENT.search(place) or DEPARTMENT.search(name)
    return f"ОМР {found.group(1)}" if found else ""


def _date(text: str) -> date | None:
    found = DATE.search(text)
    if found is None:
        return None
    day, month, year = (int(part) for part in found.groups())
    try:
        return date(year + 2000 if year < 100 else year, month, day)
    except ValueError:
        return None
