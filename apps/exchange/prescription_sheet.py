"""Разбор листа назначений врача (.docx) — TZ.md §6.1. Без Django.

Берём только то, что нужно программе: шапку, колонку «Физиотерапия» с датами и строку
консультаций из колонки «Диагностические назначения». Лекарства и анализы не читаем.

Листы врачи ведут руками в Word, поэтому разбор терпим к реальным вариантам: объединённые
ячейки, неполные строки, строка-название над заголовком таблицы, перенос списка на новую
строку. Любой странный вход даёт SheetError или предупреждение, но не падение.
"""

import re
from dataclasses import asdict, dataclass, field
from datetime import date
from typing import BinaryIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table, _Cell
from lxml.etree import XMLSyntaxError


class SheetError(ValueError):
    """Файл не похож на лист назначений. Текст показывается врачу."""


@dataclass(frozen=True)
class SheetRow:
    """Будущее назначение: исходный текст и то, что из него извлечено."""

    raw_text: str
    match_text: str
    prescribed_on: date | None = None
    cancel_date: date | None = None
    duration_min: int | None = None
    per_day: int = 1
    consultation: bool = False


@dataclass
class Sheet:
    full_name: str
    shrm: int
    room: str
    history_number: str = ""
    diagnosis: str = ""
    rows: list[SheetRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def start_date(self) -> date | None:
        """Начало курса — по физиотерапии: дата консультации, которая может и не попасть
        в программу (терапевт), не должна сдвигать курс."""
        dates = [
            row.prescribed_on for row in self.rows if row.prescribed_on and not row.consultation
        ]
        return min(dates) if dates else None

    def to_dict(self) -> dict:
        """Для временного хранения в сессии (повтор импорта без повторной загрузки файла)."""
        data = asdict(self)
        for row in data["rows"]:
            for key in ("prescribed_on", "cancel_date"):
                row[key] = row[key].isoformat() if row[key] else None
        return data

    @classmethod
    def from_dict(cls, data: dict) -> "Sheet":
        rows = []
        for row in data["rows"]:
            for key in ("prescribed_on", "cancel_date"):
                row[key] = date.fromisoformat(row[key]) if row[key] else None
            rows.append(SheetRow(**row))
        return cls(**{**data, "rows": rows})


# --- Шапка (FR-IMP-1) -------------------------------------------------------------------

HEADER_FIELDS = {
    # ФИО — до следующего поля шапки, а не до двойного пробела: двойной пробел бывает и
    # внутри ФИО («Тестова  Анна»).
    "full_name": re.compile(r"ФИО:?\s+(?P<v>.+?)\s*(?=ИБ\s*№|ШРМ|Диагноз|Палата|$)"),
    "history_number": re.compile(r"ИБ\s*№\s*(?P<v>[\w-]+)"),
    "shrm": re.compile(r"ШРМ\s*(?P<v>[345])\b"),
    "diagnosis": re.compile(r"Диагноз:\s*(?P<v>\S+)"),
    # «5а», «5 а», «9.» → «5а», «5а», «9». Палаты без номера («ВИП») — как написано.
    "room": re.compile(
        r"Палата\s*№\s*(?:(?P<num>\d+)(?:\s?(?P<letter>[а-яёa-z])(?![а-яёa-z]))?|(?P<v>\S+))",
        re.IGNORECASE,
    ),
}
FIELD_NAMES = {
    "full_name": "ФИО",
    "history_number": "ИБ",
    "shrm": "ШРМ",
    "diagnosis": "диагноз",
    "room": "палата",
}
# Длины — как у полей программы: длиннее значит, что шапка разобрана неверно.
MAX_LENGTH = {"full_name": 150, "history_number": 20, "diagnosis": 255, "room": 10}
REQUIRED = ("full_name", "shrm", "room")

# --- Таблица назначений (FR-IMP-2…5) ----------------------------------------------------

DATE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})")
DURATION = re.compile(r"(\d+)\s*мин", re.IGNORECASE)
PER_DAY = re.compile(r"(\d+)\s*р/д", re.IGNORECASE)
BULLET = re.compile(r"^\s*[-–—]\s*")
CONSULTATION = re.compile(r"^\s*консультац\w*[\s:]+", re.IGNORECASE)
GROUP_TITLE = re.compile(r"^\s*групповые\s+занятия", re.IGNORECASE)
POOL_TITLE = re.compile(r"бассейн|в\s+воде", re.IGNORECASE)
# Что остаётся от строки «30 мин 2 р/д е/д, ежедневно» без дозировки — если ничего, это не
# отдельное назначение, а дозировка для групп над ней.
DOSAGE_WORDS = re.compile(r"\d+\s*мин\w*|\d+\s*р/д|е/д|ежедневно|[\s,.;]+", re.IGNORECASE)
CONSULTATION_SPLIT = re.compile(r",|;|\s+и\s+")
HEADER_ROWS_TO_SCAN = 3
MAX_DURATION_MIN = 240
MAX_PER_DAY = 10


def parse_sheet(file: BinaryIO) -> Sheet:
    try:
        document = Document(file)
    # Переименованный .doc, архив с битым XML и т. п. — врачу нужен понятный ответ, а не 500.
    except (PackageNotFoundError, BadZipFile, XMLSyntaxError, KeyError, ValueError) as error:
        raise SheetError("Файл не читается как документ Word (.docx).") from error

    # Шапку иногда оформляют таблицей — тогда ищем её и в ячейках.
    texts = [_clean(p.text) for p in document.paragraphs]
    if not any("ФИО" in text for text in texts):
        texts += [
            _clean(p.text)
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for p in cell.paragraphs
        ]
    header = _parse_header(texts)
    found = _find_table(document.tables)
    if found is None:
        raise SheetError("В документе не найдена таблица с колонкой «Физиотерапия».")

    sheet = Sheet(**header["values"], warnings=header["warnings"])
    table, header_row = found
    sheet.rows = _parse_table(table, header_row, sheet.warnings)
    if not sheet.rows:
        sheet.warnings.append("В колонке «Физиотерапия» нет назначений.")
    return sheet


def _clean(text: str) -> str:
    return text.replace("\xa0", " ").replace("\t", " ").strip()


def _parse_header(paragraphs: list[str]) -> dict:
    found: dict[str, list[str]] = {name: [] for name in HEADER_FIELDS}
    for text in paragraphs:
        if "ФИО" not in text:
            continue
        for name, pattern in HEADER_FIELDS.items():
            match = pattern.search(text)
            if not match:
                continue
            value = _header_value(name, match)
            if value and value not in found[name]:
                found[name].append(value)

    missing = [FIELD_NAMES[name] for name in REQUIRED if not found[name]]
    if missing:
        raise SheetError(
            "В шапке листа назначений не найдено: " + ", ".join(missing) + ". "
            "Ожидается строка вида «ФИО …  ИБ №…  ШРМ 4  Диагноз: …  Палата № …»."
        )
    too_long = [
        FIELD_NAMES[name]
        for name, limit in MAX_LENGTH.items()
        if found[name] and len(found[name][0]) > limit
    ]
    if too_long:
        raise SheetError(
            "В шапке листа назначений слишком длинные значения: " + ", ".join(too_long) + ". "
            "Проверьте шапку — похоже, поля слились в одну строку."
        )

    warnings = [
        f"В листе разные значения «{FIELD_NAMES[name]}»: {', '.join(values)} — взято первое."
        for name, values in found.items()
        if len(values) > 1
    ]
    values = {name: values[0] if values else "" for name, values in found.items()}
    values["shrm"] = int(values["shrm"])
    return {"values": values, "warnings": warnings}


def _header_value(name: str, match: re.Match) -> str:
    if name == "room" and match.group("num"):
        return match.group("num") + (match.group("letter") or "").lower()
    value = match.group("v").strip().rstrip(".,;")
    return " ".join(value.split()) if name == "full_name" else value


def _find_table(tables: list[Table]) -> tuple[Table, int] | None:
    """Таблица и номер строки заголовка: над заголовком бывает строка-название."""
    for table in tables:
        for index, row in enumerate(table.rows[:HEADER_ROWS_TO_SCAN]):
            if any("Физиотерапия" in cell.text for cell in row.cells):
                return table, index
    return None


def _parse_table(table: Table, header_row: int, warnings: list[str]) -> list[SheetRow]:
    header = [_clean(cell.text) for cell in table.rows[header_row].cells]
    physio = next(i for i, text in enumerate(header) if "Физиотерапия" in text)
    diagnostics = next((i for i, text in enumerate(header) if "Диагност" in text), None)

    rows: list[SheetRow] = []
    consultations: list[SheetRow] = []
    previous: dict[int, _Cell] = {}
    for row in table.rows[header_row + 1 :]:
        cells = row.cells
        physio_cell = _cell(cells, physio, previous)
        if physio_cell is not None and CONSULTATION.match(_clean(physio_cell.text)):
            # Консультации, записанные в физиотерапию, разбираются так же, как в диагностике.
            consultations.extend(
                _consultation_rows(physio_cell.text, _date(_text(cells, physio - 1)))
            )
        elif physio_cell is not None:
            rows.extend(
                _physio_rows(
                    physio_cell.text,
                    _date(_text(cells, physio - 1)),
                    _date(_text(cells, physio + 1)),
                    warnings,
                )
            )
        if diagnostics is not None:
            diag_cell = _cell(cells, diagnostics, previous)
            if diag_cell is not None:
                consultations.extend(
                    _consultation_rows(diag_cell.text, _date(_text(cells, diagnostics - 1)))
                )
    return rows + consultations


def _cell(cells: list[_Cell], index: int, previous: dict[int, _Cell]) -> _Cell | None:
    """Ячейка колонки или None, если строка короче заголовка или это продолжение
    объединённой по вертикали ячейки (python-docx возвращает её ещё раз)."""
    if not 0 <= index < len(cells):
        return None
    cell = cells[index]
    repeated = index in previous and previous[index]._tc is cell._tc
    previous[index] = cell
    return None if repeated else cell


def _text(cells: list[_Cell], index: int) -> str:
    return cells[index].text if 0 <= index < len(cells) else ""


def _date(text: str) -> date | None:
    match = DATE.search(text)
    if not match:
        return None
    day, month, year = (int(part) for part in match.groups())
    if year < 100:
        year += 2000
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _physio_rows(
    text: str, prescribed_on: date | None, cancel_date: date | None, warnings: list[str]
) -> list[SheetRow]:
    lines = [_clean(line) for line in text.splitlines() if _clean(line)]
    if not lines:
        return []

    def row(raw: str, match: str, source: str, per_day: int | None = None) -> SheetRow:
        return SheetRow(
            raw_text=raw,
            match_text=match,
            prescribed_on=prescribed_on,
            cancel_date=cancel_date,
            duration_min=_duration(source, raw, warnings),
            per_day=per_day or _per_day(source, raw, warnings),
        )

    if not any(BULLET.match(line) or GROUP_TITLE.match(line) for line in lines):
        # Обычная ячейка: одна процедура, описание может занимать несколько строк.
        whole = " ".join(lines)
        return [row(whole, whole, whole)]

    # Ячейка со списком групп. Идём по строкам: «Групповые занятия ЛФК» (или «Бассейн»)
    # открывает список, пункты с дефисом и без — группы, строка с одной дозировкой
    # («30 мин 2 р/д») закрывает список и относится ко всем его группам. Остальное —
    # самостоятельные назначения: «Индивидуальное занятие 30 мин», «Массаж спины».
    # «2 р/д» под списком — сколько групповых занятий в день на весь список, т. е. каждая
    # группа идёт один раз (решение владельца 04.10, TZ.md решение 42).
    result: list[SheetRow] = []
    title, groups, in_list = "", [], False

    def flush(dosage: str) -> None:
        total = _int(PER_DAY, dosage)
        if total is not None and groups and total != len(groups):
            names = ", ".join(groups)
            warnings.append(
                f"«{title or 'Группы'}: {names}»: {total} р/д на {_groups(len(groups))} — "
                "каждая группа поставлена 1 раз в день. Если группа нужна чаще, "
                "поправьте частоту в назначении."
            )
        for group in groups:
            # «Бассейн / - верхняя конечность»: признак бассейна — в заголовке списка.
            match = f"{title} {group}" if POOL_TITLE.search(title) else group
            raw = f"{title}: {group}" if title else group
            result.append(row(raw, match, f"{group} {dosage}", per_day=1))
        groups.clear()

    for line in lines:
        if BULLET.match(line):
            groups.append(BULLET.sub("", line))
            in_list = True
        elif GROUP_TITLE.match(line) or (POOL_TITLE.search(line) and not _has_dosage(line)):
            flush("")
            title, in_list = line, True
        elif _is_dosage_only(line):
            if groups:
                flush(line)
                in_list = False
            elif result:
                # Дозировка на отдельной строке под процедурой — относится к ней.
                last = result.pop()
                result.append(row(last.raw_text, last.match_text, f"{last.raw_text} {line}"))
        elif in_list and not _has_dosage(line):
            groups.append(line)
        else:
            flush("")
            title, in_list = "", False
            result.append(row(line, line, line))
    flush("")
    return result


def _is_dosage_only(line: str) -> bool:
    return _has_dosage(line) and not DOSAGE_WORDS.sub("", line)


def _has_dosage(line: str) -> bool:
    return bool(DURATION.search(line) or PER_DAY.search(line))


def _duration(text: str, label: str, warnings: list[str]) -> int | None:
    value = _int(DURATION, text)
    if value is not None and not 1 <= value <= MAX_DURATION_MIN:
        warnings.append(f"«{label}»: длительность {value} мин не похожа на правду — не учтена.")
        return None
    return value


def _groups(count: int) -> str:
    if count % 10 == 1 and count % 100 != 11:
        return f"{count} группу"
    if count % 10 in (2, 3, 4) and count % 100 not in (12, 13, 14):
        return f"{count} группы"
    return f"{count} групп"


def _per_day(text: str, label: str, warnings: list[str]) -> int:
    value = _int(PER_DAY, text)
    if value is not None and not 1 <= value <= MAX_PER_DAY:
        warnings.append(f"«{label}»: частота {value} р/д не похожа на правду — поставлено 1.")
        return 1
    return value or 1


def _consultation_rows(text: str, prescribed_on: date | None) -> list[SheetRow]:
    """Строка «Консультация …» вместе с продолжением на следующих строках ячейки:
    продолжение — строка после запятой или начинающаяся со строчной буквы."""
    lines = [_clean(line) for line in text.splitlines() if _clean(line)]
    rows = []
    index = 0
    while index < len(lines):
        if not CONSULTATION.match(lines[index]):
            index += 1
            continue
        parts = [CONSULTATION.sub("", lines[index])]
        index += 1
        while index < len(lines) and (parts[-1].endswith(",") or lines[index][:1].islower()):
            parts.append(lines[index])
            index += 1
        for part in CONSULTATION_SPLIT.split(" ".join(parts).rstrip(".")):
            part = part.strip()
            if part:
                rows.append(
                    SheetRow(
                        raw_text=f"Консультация {part}",
                        match_text=part,
                        prescribed_on=prescribed_on,
                        consultation=True,
                    )
                )
    return rows


def _int(pattern: re.Pattern, text: str) -> int | None:
    match = pattern.search(text)
    return int(match.group(1)) if match else None
