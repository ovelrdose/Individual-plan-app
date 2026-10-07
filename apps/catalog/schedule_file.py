"""Чтение расписания групп из xlsx заказчика (lfk.xlsx, basseyn.xlsx, Gruppy_DS.xlsx).

Формат: колонка A — время начала, колонка B — название группы, колонка C (необязательная) —
место проведения («181 кабинет»), без заголовка.
В фазе 3 этот разбор используется и для импорта через веб (FR-EXC-1).
"""

import re
from dataclasses import dataclass
from datetime import time
from pathlib import Path

from openpyxl import load_workbook

from .domain import excel_time


@dataclass(frozen=True)
class ScheduleRow:
    start: time
    name: str
    cell: str
    place: str = ""


# «179кабинет» в файле — «179 кабинет».
_NUMBER_WORD = re.compile(r"(\d)(?=[^\d\s])")


class ScheduleFileError(ValueError):
    pass


def read_schedule(path: Path) -> list[ScheduleRow]:
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = []
    for number, values in enumerate(sheet.iter_rows(max_col=3, values_only=True), start=1):
        start, name, place = [*values, None, None, None][:3]
        if start is None and name is None:
            continue
        if start is None or not isinstance(name, str) or not name.strip():
            raise ScheduleFileError(f"{path.name}, строка {number}: нужны время и название.")
        try:
            parsed = excel_time(start)
        except ValueError as error:
            raise ScheduleFileError(f"{path.name}, ячейка A{number}: {error}") from error
        rows.append(ScheduleRow(parsed, name.strip(), f"A{number}", _place(place)))
    return rows


def _place(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(_NUMBER_WORD.sub(r"\1 ", value).split())
