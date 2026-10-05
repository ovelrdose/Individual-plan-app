"""Выгрузка шахматки на дату в .xlsx (TZ.md FR-CRD-7) — чистый Python без Django.

A1 — дата, строка 1 — фамилии (пары «A/ B»), колонка A — интервалы слотов, справа через
пустую колонку — тренажёры. Альбомная ориентация, вписать в одну страницу по ширине.
Дата и время тренажёров — настоящие значения Excel.
"""

from dataclasses import dataclass, field
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DATE_FORMAT = "dd.mm.yyyy"
TIME_FORMAT = "h:mm"
BUSY_FILL = PatternFill(fill_type="solid", start_color="FFD9D9D9", end_color="FFD9D9D9")
OVER_FILL = PatternFill(fill_type="solid", start_color="FFFFF2CC", end_color="FFFFF2CC")
THIN = Side(style="thin", color="FF808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass(frozen=True)
class BoardCell:
    text: str = ""
    busy: bool = False  # блок — серым, как в бумажной шахматке


@dataclass(frozen=True)
class BoardSheet:
    """Шахматка на дату: ``cells[строка слота][колонка]``, ``equipment_cells[строка
    времени][тренажёр]`` — список записей и «сверх вместимости»."""

    day: date
    slots: list[tuple[time, time]]
    columns: list[str]
    cells: list[list[BoardCell]]
    equipment: list[str] = field(default_factory=list)
    equipment_times: list[time] = field(default_factory=list)
    equipment_cells: list[list[tuple[list[str], bool]]] = field(default_factory=list)


def render_board(data: BoardSheet) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{data.day:%d.%m}"

    corner = sheet.cell(1, 1, datetime.combine(data.day, time()))
    corner.number_format = DATE_FORMAT
    corner.font = Font(bold=True)
    for index, label in enumerate(data.columns, start=2):
        _head(sheet.cell(1, index, label))
    for row, (start, end) in enumerate(data.slots, start=2):
        _head(sheet.cell(row, 1, f"{start.hour}:{start.minute:02d}-{end.hour}:{end.minute:02d}"))
        for column, cell in enumerate(data.cells[row - 2], start=2):
            target = sheet.cell(row, column, cell.text or None)
            target.border = BORDER
            target.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.busy:
                target.fill = BUSY_FILL
    sheet.column_dimensions["A"].width = 12
    for index in range(2, len(data.columns) + 2):
        sheet.column_dimensions[get_column_letter(index)].width = 16

    if data.equipment:
        first = len(data.columns) + 3  # через пустую колонку
        _head(sheet.cell(1, first, "Время"))
        for offset, name in enumerate(data.equipment, start=1):
            _head(sheet.cell(1, first + offset, name))
            sheet.column_dimensions[get_column_letter(first + offset)].width = 18
        sheet.column_dimensions[get_column_letter(first)].width = 8
        for row, start in enumerate(data.equipment_times, start=2):
            moment = sheet.cell(row, first, start)
            moment.number_format = TIME_FORMAT
            _head(moment)
            for offset, (names, over) in enumerate(data.equipment_cells[row - 2], start=1):
                target = sheet.cell(row, first + offset, "\n".join(names) or None)
                target.border = BORDER
                target.alignment = Alignment(wrap_text=True, vertical="top")
                if over:
                    target.fill = OVER_FILL

    sheet.freeze_panes = "B2"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _head(cell) -> None:
    cell.font = Font(bold=True)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
