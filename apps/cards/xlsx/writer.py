"""Заполнение шаблона карты (TZ.md, FR-CRD-2).

Оформление шаблона не трогаем: пишем только значения, форматы чисел и серую заливку.
"""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .data import CardData
from .formatting import sex_age_text
from .mapping import CardMapping

TIME_FORMAT = "h:mm"
DATE_FORMAT = "dd.mm"
INACTIVE_FILL = PatternFill(fill_type="solid", start_color="FFD9D9D9", end_color="FFD9D9D9")


class CardRenderError(ValueError):
    """Данные не помещаются в шаблон. Сообщение показывается пользователю."""


def render_card(template_path: Path, mapping: CardMapping, data: CardData) -> bytes:
    _check_fits(mapping, data)

    workbook = load_workbook(template_path)
    sheet = workbook.active
    _fill_header(sheet, mapping, data)
    _fill_schedule(sheet, mapping, data)
    _fill_date_blocks(sheet, mapping, data)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _check_fits(mapping: CardMapping, data: CardData) -> None:
    errors = []
    schedule_capacity = len(mapping.schedule.rows)
    if len(data.schedule) > schedule_capacity:
        labels = ", ".join(f"{item.start:%H:%M} {item.label or 'окно'}" for item in data.schedule)
        errors.append(
            f"В расписании {len(data.schedule)} строк (слоты сетки инструкторов), а в карте "
            f"помещается {schedule_capacity}: {labels}. Проверьте дневные слоты в справочнике."
        )
    if len(data.procedures) > mapping.max_procedures:
        errors.append(
            f"Процедур для карты {len(data.procedures)}, а в карте помещается "
            f"{mapping.max_procedures}."
        )
    if len(data.course_dates) > mapping.max_dates:
        errors.append(
            f"Дней курса {len(data.course_dates)}, а в карте помещается {mapping.max_dates}."
        )
    if errors:
        raise CardRenderError(" ".join(errors))


def _fill_header(sheet: Worksheet, mapping: CardMapping, data: CardData) -> None:
    sheet[mapping.department] = data.department
    sheet[mapping.full_name] = data.full_name
    parts = [sex_age_text(data.sex, data.age)]
    if data.withdrawn_on is not None:
        parts.append(f"выбыл {data.withdrawn_on:%d.%m}")
    sheet[mapping.sex_age] = ", ".join(part for part in parts if part)
    sheet[mapping.doctor] = data.doctor


def _fill_schedule(sheet: Worksheet, mapping: CardMapping, data: CardData) -> None:
    """Строки идут в порядке сетки; у окна есть только время — в карте видно, когда
    пациент свободен (TZ.md, FR-CRD-2)."""
    area = mapping.schedule
    for row, item in zip(area.rows, data.schedule, strict=False):
        time_cell = sheet[f"{area.time_col}{row}"]
        time_cell.value = item.start
        time_cell.number_format = TIME_FORMAT
        sheet[f"{area.label_col}{row}"] = item.label or None
        sheet[f"{area.place_col}{row}"] = item.place or None


def _fill_date_blocks(sheet: Worksheet, mapping: CardMapping, data: CardData) -> None:
    remaining = list(data.course_dates)
    for block in mapping.date_blocks:
        block_dates, remaining = remaining[: block.size], remaining[block.size :]
        for column, day in zip(block.date_columns, block_dates, strict=False):
            cell = sheet[f"{column}{block.date_row}"]
            cell.value = day
            cell.number_format = DATE_FORMAT

        for row, procedure in zip(block.procedure_rows, data.procedures, strict=False):
            sheet[f"{mapping.procedure_label_col}{row}"] = procedure.label
            for column, day in zip(block.date_columns, block_dates, strict=False):
                if day in data.rest_dates or not procedure.is_active(day):
                    sheet[f"{column}{row}"].fill = INACTIVE_FILL
