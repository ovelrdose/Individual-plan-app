"""Подготовка шаблонов карт из файла заказчика (TZ.md, FR-CRD-1).

Из `programm_template.xlsx` каждый лист «ШРМ N» сохраняется отдельным файлом,
а примерные данные (пациент, врач, расписание, даты, процедуры) стираются по mapping.
Карта печатается на двух сторонах листа А4 (TZ.md, решение 74): лицевая — расписание, логотип,
пациент, врач и «Уважаемый пациент!»; оборотная — таблица дат.
"""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from .mapping import CardMapping

SOURCE_SHEET = "ШРМ {shrm}"
TEMPLATE_FILE = "shrm{shrm}.xlsx"
# Средняя ширина буквы Calibri в долях кегля, с запасом: обычный и полужирный.
CHAR_EM, BOLD_CHAR_EM = 0.5, 0.55
DEFAULT_COLUMN_WIDTH = 8.43
# A4 альбомной ориентации в пунктах.
PAGE_WIDTH, PAGE_HEIGHT = 842, 595
# Запас на расхождение оценки ширины с отрисовкой Excel (шрифты, принтер).
SCALE_RESERVE = 0.92


def extract_templates(
    source: Path, mappings: dict[int, CardMapping], target_dir: Path
) -> list[Path]:
    target_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for shrm, mapping in sorted(mappings.items()):
        workbook = load_workbook(source)
        keep = SOURCE_SHEET.format(shrm=shrm)
        for sheet in list(workbook.worksheets):
            if sheet.title != keep:
                workbook.remove(sheet)
        workbook.active = 0
        # В файле заказчика стоит ручной пересчёт — он «прилипает» к сеансу Excel пользователя.
        workbook.calculation.calcMode = "auto"
        clear_data(workbook.active, mapping)
        set_print_layout(workbook.active, mapping)

        path = target_dir / TEMPLATE_FILE.format(shrm=shrm)
        workbook.save(path)
        written.append(path)
    return written


def data_cells(mapping: CardMapping) -> list[str]:
    """Все ячейки, которые заполняет выгрузка, — в шаблоне они должны быть пустыми."""
    cells = [mapping.department, mapping.full_name, mapping.room, mapping.doctor]
    area = mapping.schedule
    for row in area.rows:
        cells += [f"{col}{row}" for col in _columns_between(area.time_col, area.place_col)]
    for block in mapping.date_blocks:
        cells += [f"{col}{block.date_row}" for col in block.date_columns]
        last_col = block.date_columns[-1]
        for row in block.procedure_rows:
            cells += [
                f"{col}{row}" for col in _columns_between(mapping.procedure_label_col, last_col)
            ]
    return cells


def clear_data(sheet: Worksheet, mapping: CardMapping) -> None:
    for coordinate in data_cells(mapping):
        cell = sheet[coordinate]
        if not isinstance(cell, MergedCell):
            cell.value = None


def set_print_layout(sheet: Worksheet, mapping: CardMapping) -> None:
    """Область печати — от A1 до правого нижнего края карты, две страницы А4: разрыв перед
    таблицей дат. «Вписать в N страниц» Excel применяет без ручных разрывов, поэтому масштаб
    задаётся числом — наибольший, при котором каждая сторона помещается на лист."""
    last_col, last_row = _card_bounds(sheet, mapping)
    front_last_row = mapping.date_blocks[0].date_row - 1
    sheet.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    sheet.row_breaks = RowBreak()
    sheet.row_breaks.append(Break(id=front_last_row))
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
    sheet.page_setup.fitToWidth = None
    sheet.page_setup.fitToHeight = None
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = print_scale(
        sheet, last_col, [(1, front_last_row), (front_last_row + 1, last_row)]
    )
    sheet.print_options.horizontalCentered = True


def print_scale(sheet: Worksheet, last_col: int, pages: list[tuple[int, int]]) -> int:
    """Масштаб в процентах, при котором колонки A…last_col и каждая группа строк
    помещаются на страницу А4 альбомной ориентации с полями шаблона."""
    margins = sheet.page_margins
    width = PAGE_WIDTH - (margins.left + margins.right) * 72
    height = PAGE_HEIGHT - (margins.top + margins.bottom) * 72
    content_width = sum(_column_points(sheet, col) for col in range(1, last_col + 1))
    content_height = max(
        sum(_row_points(sheet, row) for row in range(first, last + 1)) for first, last in pages
    )
    scale = min(width / content_width, height / content_height) * SCALE_RESERVE
    return max(10, min(100, int(scale * 100)))


def column_width(sheet: Worksheet, index: int) -> float:
    """Ширина колонки в символах. Не через ``sheet.column_dimensions[буква]``: openpyxl
    при таком обращении создаёт запись шириной 13, и она перекрывает диапазон колонок."""
    for dim in sheet.column_dimensions.values():
        if dim.min is not None and dim.min <= index <= dim.max:
            return 0.0 if dim.hidden else (dim.width or DEFAULT_COLUMN_WIDTH)
    return sheet.sheet_format.defaultColWidth or DEFAULT_COLUMN_WIDTH


def _column_points(sheet: Worksheet, index: int) -> float:
    # Ширина в пикселях для Calibri 11 (символ — 7 px, плюс 5 px полей), пиксель — 0,75 пт.
    width = column_width(sheet, index)
    return int(width * 7 + 5) * 0.75 if width else 0.0


def _row_points(sheet: Worksheet, row: int) -> float:
    dim = sheet.row_dimensions.get(row)
    if dim is not None and dim.hidden:
        return 0.0
    if dim is not None and dim.height is not None:
        return dim.height
    return sheet.sheet_format.defaultRowHeight or 15.0


def _card_bounds(sheet: Worksheet, mapping: CardMapping) -> tuple[int, int]:
    """Край карты: ячейки с текстом или рамкой, логотипы и ячейки, которые заполнит выгрузка."""
    last_col = last_row = 1
    for row in sheet.iter_rows():
        for cell in row:
            border = cell.border
            has_border = any(
                getattr(border, side).style for side in ("left", "right", "top", "bottom")
            )
            if cell.value is not None or has_border:
                last_col, last_row = max(last_col, cell.column), max(last_row, cell.row)
            if isinstance(cell.value, str):
                last_col = max(last_col, _text_end_column(sheet, cell))
    for image in sheet._images:
        # Индексы якоря считаются с нуля.
        last_col = max(last_col, image.anchor.to.col + 1)
        last_row = max(last_row, image.anchor.to.row + 1)
    for coordinate in data_cells(mapping):
        cell = sheet[coordinate]
        last_col, last_row = max(last_col, cell.column), max(last_row, cell.row)
    return last_col, last_row


def _text_end_column(sheet: Worksheet, cell) -> int:
    """Колонка, до которой Excel дотягивает текст, вылезающий из ячейки вправо.

    Оценка грубая (буква Calibri ≈ полкегля), но достаточная, чтобы область печати
    не обрезала заголовки.
    """
    merged = any(cell.coordinate in rng for rng in sheet.merged_cells.ranges)
    if (
        merged
        or cell.alignment.wrap_text
        or cell.alignment.horizontal not in (None, "general", "left")
    ):
        return cell.column
    em = BOLD_CHAR_EM if cell.font.b else CHAR_EM
    need = len(cell.value.rstrip()) * em * (cell.font.sz or 11)
    column = cell.column
    while True:
        need -= _column_points(sheet, column)
        if need <= 0:
            return column
        column += 1


def _columns_between(first: str, last: str) -> list[str]:
    start, end = column_index_from_string(first), column_index_from_string(last)
    return [get_column_letter(index) for index in range(start, end + 1)]
