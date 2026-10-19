"""Генератор обезличенного листа назначений той же структуры, что primary_docs/pacient.docx."""

import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from docx import Document

HEADER = (
    "ФИО  {fio}      ИБ №{ib}      ШРМ {shrm}          Диагноз: {dx}\xa0"
    "                Палата № {room}"
)
COLUMNS = [
    "Дата назначения",
    "Диагностические назначения",
    "Дата",
    "Физиотерапия",
    "Дата отмены",
    "Дата назначения",
    "Наружние назначения.",
]


def make_sheet(
    *,
    fio: str = "Тестова Анна Сергеевна",
    ib: str = "5402",
    shrm: int | str = 4,
    dx: str = "Т91.3",
    room: str = "5а",
    rows: list[tuple[str, str, str]] | None = None,
    diagnostics: list[str] | None = None,
    second_header: dict | None = None,
    header: str | None = None,
    with_table: bool = True,
) -> BytesIO:
    """rows — (дата, текст физиотерапии, дата отмены); diagnostics — тексты диагностики."""
    rows = rows if rows is not None else [("02.10.26", "Имитрон  15 мин 1 р/д е/д", "")]
    document = Document()
    document.add_paragraph("Лист назначений")
    values = {"fio": fio, "ib": ib, "shrm": shrm, "dx": dx, "room": room}
    document.add_paragraph(header if header is not None else HEADER.format(**values))
    if with_table:
        table = document.add_table(rows=1, cols=len(COLUMNS))
        for cell, title in zip(table.rows[0].cells, COLUMNS, strict=True):
            cell.text = title
        diagnostics = diagnostics or []
        for index in range(max(len(rows), len(diagnostics))):
            cells = table.add_row().cells
            if index < len(rows):
                prescribed, text, cancel = rows[index]
                cells[2].text, cells[3].text, cells[4].text = prescribed, text, cancel
            if index < len(diagnostics):
                cells[0].text, cells[1].text = "02.10.26", diagnostics[index]
    if second_header is not None:
        document.add_paragraph("Лист назначений")
        document.add_paragraph(HEADER.format(**(values | second_header)))
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    buffer.name = "лист.docx"
    return buffer


def make_doc_sheet(**kwargs) -> BytesIO:
    """Тот же лист в старом формате Word (.doc) — через LibreOffice; без него тест пропускается.

    Сразу в .doc таблица python-docx теряется (у неё нет сетки колонок), поэтому через .odt —
    так получается файл, как из настоящего Word.
    """
    import pytest

    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("Нет LibreOffice (soffice) — .doc не создать")
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        (work / "sheet.docx").write_bytes(make_sheet(**kwargs).getvalue())
        profile = f"-env:UserInstallation={(work / 'profile').as_uri()}"
        for source, target in (("sheet.docx", "odt"), ("sheet.odt", "doc")):
            subprocess.run(
                [soffice, profile, "--headless", "--convert-to", target, source],
                cwd=work,
                capture_output=True,
                timeout=60,
                check=True,
            )
        buffer = BytesIO((work / "sheet.doc").read_bytes())
    buffer.name = "лист.doc"
    return buffer
