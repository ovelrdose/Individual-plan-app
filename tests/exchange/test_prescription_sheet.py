from datetime import date
from io import BytesIO

import pytest
from docx import Document

from apps.exchange.prescription_sheet import Sheet, SheetError, parse_sheet
from apps.programs.domain import SHRM_UK
from tests.sheets import COLUMNS, HEADER, make_sheet

HEADER_TEXT = HEADER.format(fio="Тестова Анна Сергеевна", ib="1", shrm=4, dx="I69", room="5")
PHYSIO = 3  # колонка make_sheet


def header(fio: str = "Тестова Анна Сергеевна", rest: str = "ИБ №5402  ШРМ 4  Палата № 5а") -> str:
    return f"ФИО  {fio}      {rest}"


def document(build) -> BytesIO:
    """Лист с таблицей, которую собирает build(table) — для вариантов, которых нет в make_sheet."""
    doc = Document()
    doc.add_paragraph(HEADER_TEXT)
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    for cell, title in zip(table.rows[0].cells, COLUMNS, strict=True):
        cell.text = title
    build(doc, table)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def edited(buffer: BytesIO, change) -> BytesIO:
    """Лист make_sheet, изменённый функцией change(document) — объединения ячеек."""
    document = Document(buffer)
    change(document)
    result = BytesIO()
    document.save(result)
    result.seek(0)
    return result


def physio(sheet):
    return [r.raw_text for r in sheet.rows if not r.consultation]


GROUPS = "Групповые занятия ЛФК\n- Эрго общая\n– I can нога\n30 мин 2 р/д е/д"


class TestHeader:
    def test_fields(self):
        sheet = parse_sheet(
            make_sheet(fio="Тестова Анна Сергеевна", ib="77", shrm=5, dx="I69.4", room="5а")
        )

        assert (sheet.full_name, sheet.history_number, sheet.shrm, sheet.room) == (
            "Тестова Анна Сергеевна",
            "77",
            5,
            "5а",
        )
        assert sheet.warnings == []
        # Диагноз в систему не переносится (TZ.md, решение 70) — даже во временную сессию.
        assert "I69.4" not in str(sheet.to_dict())

    def test_different_values_in_repeated_header(self):
        sheet = parse_sheet(make_sheet(ib="5402", second_header={"ib": "513402", "room": "7"}))

        assert sheet.history_number == "5402"
        assert sheet.room == "5а"
        assert "«ИБ»: 5402, 513402" in sheet.warnings[0]
        assert "«палата»: 5а, 7" in sheet.warnings[1]

    def test_same_repeated_header_is_not_a_warning(self):
        assert parse_sheet(make_sheet(second_header={})).warnings == []

    def test_optional_fields(self):
        sheet = parse_sheet(
            make_sheet(header="ФИО  Тестов Тест Тестович      ШРМ 3          Палата № 9")
        )

        assert (sheet.full_name, sheet.history_number) == ("Тестов Тест Тестович", "")

    @pytest.mark.parametrize(
        "rest",
        [
            "ИБ №5402  ШРМ УК2  Палата № 5а",
            "ИБ №5402  ШРМ: УК-3  Палата № 5а",
            "ИБ №5402  ШРМ УК  Палата № 5а",
            "УК 2  ИБ №5402  Палата № 5а",
        ],
    )
    def test_uk_instead_of_shrm(self, rest):
        sheet = parse_sheet(make_sheet(header=header(rest=rest)))

        assert (sheet.full_name, sheet.shrm, sheet.room) == (
            "Тестова Анна Сергеевна",
            SHRM_UK,
            "5а",
        )

    def test_uk_inside_surname_is_not_uk(self):
        with pytest.raises(SheetError, match="ШРМ"):
            parse_sheet(make_sheet(header=header(fio="КУКУШКИНА Анна", rest="Палата № 5а")))

    def test_missing_required_fields(self):
        with pytest.raises(SheetError, match="ШРМ, палата"):
            parse_sheet(make_sheet(header="ФИО  Тестов Тест Тестович      ИБ №5"))

    def test_no_header(self):
        with pytest.raises(SheetError, match="ФИО"):
            parse_sheet(make_sheet(header="Просто текст"))


class TestTable:
    def test_groups_become_separate_rows(self):
        sheet = parse_sheet(make_sheet(rows=[("02.10.26", GROUPS, "")]))

        assert [(r.raw_text, r.match_text, r.duration_min, r.per_day) for r in sheet.rows] == [
            ("Групповые занятия ЛФК: Эрго общая", "Эрго общая", 30, 1),
            ("Групповые занятия ЛФК: I can нога", "I can нога", 30, 1),
        ]
        # «2 р/д» на две группы — каждая раз в день, предупреждать не о чем (решение 42).
        assert sheet.warnings == []

    def test_group_list_frequency_not_matching_groups_is_a_warning(self):
        text = "Групповые занятия ЛФК\n- Эрго общая\n- I can нога\n- Нейро-тренинг\n30 мин 2 р/д"
        sheet = parse_sheet(make_sheet(rows=[("02.10.26", text, "")]))

        assert [r.per_day for r in sheet.rows] == [1, 1, 1]
        (warning,) = sheet.warnings
        assert "2 р/д на 3 группы — каждая группа поставлена 1 раз в день" in warning

    @pytest.mark.parametrize(
        ("text", "duration", "per_day"),
        [
            ("Индивидуальное занятие ЛФК 30 мин 1 р/д е/д", 30, 1),
            ("Степпер 15 минут 1 р/д, ежедневно", 15, 1),
            ("Thera Trainer Tigo 15 мин 2 р/д е/д", 15, 2),
            ("Массаж", None, 1),
        ],
    )
    def test_duration_and_frequency(self, text, duration, per_day):
        (row,) = parse_sheet(make_sheet(rows=[("02.10.26", text, "")])).rows

        assert (row.raw_text, row.duration_min, row.per_day) == (text, duration, per_day)

    def test_dates(self):
        sheet = parse_sheet(
            make_sheet(
                rows=[
                    ("03.10.26", "Имитрон 15 мин", ""),
                    ("02.10.2026", "Степпер 15 мин", "09.10.26"),
                    ("", "Алмаг 20 мин", "31.02.26"),  # несуществующая дата — игнорируется
                ]
            )
        )

        assert [(r.prescribed_on, r.cancel_date) for r in sheet.rows] == [
            (date(2026, 10, 3), None),
            (date(2026, 10, 2), date(2026, 10, 9)),
            (None, None),
        ]
        assert sheet.start_date == date(2026, 10, 2)

    def test_empty_rows_are_skipped(self):
        sheet = parse_sheet(
            make_sheet(rows=[("", "", ""), ("02.10.26", "Имитрон", ""), ("", " \n ", "")])
        )

        assert [r.raw_text for r in sheet.rows] == ["Имитрон"]

    def test_no_physio_rows_is_a_warning(self):
        sheet = parse_sheet(make_sheet(rows=[]))

        assert sheet.rows == []
        assert sheet.start_date is None
        assert "нет назначений" in sheet.warnings[0]

    def test_consultations(self):
        sheet = parse_sheet(
            make_sheet(
                rows=[("02.10.26", "Имитрон", "")],
                diagnostics=[
                    "Консультация терапевта, медицинского психолога, медицинского логопеда, "
                    "эрготерапевта и физического терапевта.",
                    "Гликемический профиль на 03.10.26",
                ],
            )
        )
        consultations = [r for r in sheet.rows if r.consultation]

        assert [r.match_text for r in consultations] == [
            "терапевта",
            "медицинского психолога",
            "медицинского логопеда",
            "эрготерапевта",
            "физического терапевта",
        ]
        assert consultations[1].raw_text == "Консультация медицинского психолога"
        assert sheet.rows[0].raw_text == "Имитрон", "физиотерапия идёт раньше консультаций"

    def test_no_physio_table(self):
        with pytest.raises(SheetError, match="Физиотерапия"):
            parse_sheet(make_sheet(with_table=False))

    def test_not_a_docx(self):
        with pytest.raises(SheetError, match=r"\.docx"):
            parse_sheet(BytesIO(b"definitely not a zip"))


def test_session_roundtrip():
    sheet = parse_sheet(
        make_sheet(
            rows=[("02.10.26", "Имитрон", "05.10.26")], diagnostics=["Консультация психолога"]
        )
    )

    restored = Sheet.from_dict(sheet.to_dict())

    assert restored == sheet


def test_zip_with_broken_document():
    import zipfile

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<broken")
    buffer.seek(0)

    with pytest.raises(SheetError):
        parse_sheet(buffer)


# --- Варианты реальных листов (TZ.md, решение 24) -------------------------------------


class TestHeaderVariants:
    @pytest.mark.parametrize(
        "text",
        [
            "ФИО\tТестова Анна Сергеевна\tИБ №5402\tШРМ 4\tДиагноз: Т91.3\tПалата № 5а",
            "ФИО\xa0\xa0Тестова Анна Сергеевна\xa0\xa0ИБ\xa0№5402\xa0ШРМ\xa04\xa0"
            "Диагноз:\xa0Т91.3\xa0Палата\xa0№\xa05а",
            "ФИО Тестова Анна Сергеевна ИБ № 5402 ШРМ 4 Диагноз: Т91.3 Палата № 5а",
        ],
        ids=["tabs", "nbsp", "single-spaces"],
    )
    def test_separators(self, text):
        sheet = parse_sheet(make_sheet(header=text))

        assert (sheet.full_name, sheet.history_number, sheet.shrm, sheet.room) == (
            "Тестова Анна Сергеевна",
            "5402",
            4,
            "5а",
        )
        assert sheet.warnings == []

    def test_fields_in_other_order(self):
        sheet = parse_sheet(
            make_sheet(header="ФИО  Тестов Тест Тестович  Палата № 12  ИБ № 77  ШРМ 5")
        )

        assert (sheet.full_name, sheet.history_number, sheet.shrm, sheet.room) == (
            "Тестов Тест Тестович",
            "77",
            5,
            "12",
        )

    def test_double_surname(self):
        sheet = parse_sheet(make_sheet(fio="Тестова-Пробная Анна Сергеевна"))

        assert sheet.full_name == "Тестова-Пробная Анна Сергеевна"

    def test_double_space_inside_full_name(self):
        sheet = parse_sheet(make_sheet(header=header(fio="Тестова  Анна Сергеевна")))

        assert sheet.full_name.split() == ["Тестова", "Анна", "Сергеевна"]

    @pytest.mark.parametrize(
        ("room", "expected"),
        [("5а", "5а"), ("5 а", "5а"), ("5А", "5а"), ("9.", "9"), ("12", "12"), ("ВИП", "ВИП")],
    )
    def test_room(self, room, expected):
        assert parse_sheet(make_sheet(room=room)).room == expected

    def test_room_followed_by_word(self):
        header = "ФИО  Тестова Анна Сергеевна      ШРМ 4      Палата № 5 Диагноз: I69"

        assert parse_sheet(make_sheet(header=header)).room == "5"

    @pytest.mark.parametrize("kwargs", [{"room": "12345678901234"}, {"ib": "1" * 25}])
    def test_too_long_values(self, kwargs):
        with pytest.raises(SheetError, match="слишком длинные"):
            parse_sheet(make_sheet(**kwargs))


class TestTableVariants:
    @pytest.mark.parametrize(
        ("prescribed", "cancel"),
        [("2.10.26", "5.10.26"), ("02.10.2026", "05.10.2026"), ("2.10.2026", "05.10.26")],
    )
    def test_date_formats(self, prescribed, cancel):
        (row,) = parse_sheet(make_sheet(rows=[(prescribed, "Имитрон", cancel)])).rows

        assert (row.prescribed_on, row.cancel_date) == (date(2026, 10, 2), date(2026, 10, 5))

    def test_multi_paragraph_cell_with_blank_lines_and_nbsp(self):
        text = "\nИмитрон\xa0\xa015\xa0мин\n\n \n1\xa0р/д е/д\n"

        (row,) = parse_sheet(make_sheet(rows=[("02.10.26", text, "")])).rows

        assert (row.raw_text.split(), row.duration_min, row.per_day) == (
            ["Имитрон", "15", "мин", "1", "р/д", "е/д"],
            15,
            1,
        )

    def test_bullets_with_tab_and_em_dash(self):
        text = "Групповые занятия ЛФК\n-\tЭрго общая\n—  Нейро - тренинг\n30 мин 1 р/д"

        rows = parse_sheet(make_sheet(rows=[("02.10.26", text, "")])).rows

        assert [r.match_text for r in rows] == ["Эрго общая", "Нейро - тренинг"]

    def test_two_group_blocks_in_one_cell(self):
        text = (
            "Групповые занятия ЛФК\n- Эрго общая\n30 мин 2 р/д\n"
            "Групповые занятия ЛФК\n- Нейро-тренинг\n15 мин 1 р/д"
        )

        rows = parse_sheet(make_sheet(rows=[("02.10.26", text, "")])).rows

        assert [(r.match_text, r.duration_min, r.per_day) for r in rows] == [
            ("Эрго общая", 30, 1),
            ("Нейро-тренинг", 15, 1),
        ]

    def test_extra_line_after_group_list(self):
        sheet = parse_sheet(
            make_sheet(
                rows=[
                    (
                        "02.10.26",
                        "Групповые занятия ЛФК\n- Эрго общая\n30 мин 2 р/д\nМассаж спины",
                        "",
                    )
                ]
            )
        )

        assert physio(sheet) == ["Групповые занятия ЛФК: Эрго общая", "Массаж спины"]
        assert (sheet.rows[0].duration_min, sheet.rows[0].per_day) == (30, 1)
        assert (sheet.rows[1].duration_min, sheet.rows[1].per_day) == (None, 1)
        assert "2 р/д на 1 группу" in sheet.warnings[0]

    def test_consultations_with_and_and_dot(self):
        rows = parse_sheet(
            make_sheet(diagnostics=["Консультация медицинского психолога и логопеда."])
        ).rows

        assert [r.match_text for r in rows if r.consultation] == [
            "медицинского психолога",
            "логопеда",
        ]

    def test_consultations_with_colon(self):
        rows = parse_sheet(make_sheet(diagnostics=["Консультации: психолога, логопеда"])).rows

        assert [r.match_text for r in rows if r.consultation] == ["психолога", "логопеда"]

    def test_consultation_list_continues_on_next_line(self):
        sheet = parse_sheet(
            make_sheet(
                diagnostics=[
                    "Консультация терапевта, медицинского психолога,\nлогопеда, эрготерапевта\n"
                    "Гликемический профиль"
                ]
            )
        )

        assert [r.match_text for r in sheet.rows if r.consultation] == [
            "терапевта",
            "медицинского психолога",
            "логопеда",
            "эрготерапевта",
        ]

    def test_vertically_merged_physio_cell_is_read_once(self):
        def build(doc, table):
            first, second = table.add_row().cells, table.add_row().cells
            first[2].text, first[3].text = "02.10.26", "Имитрон 15 мин 1 р/д"
            first[3].merge(second[3])

        assert physio(parse_sheet(document(build))) == ["Имитрон 15 мин 1 р/д"]

    def test_horizontally_merged_header_does_not_crash(self):
        def merge(document):
            table = document.tables[0]
            table.cell(0, 0).merge(table.cell(0, 1))

        buffer = edited(make_sheet(rows=[("02.10.26", "Имитрон", "05.10.26")]), merge)

        rows = parse_sheet(buffer).rows
        assert rows[0].raw_text == "Имитрон"
        assert rows[0].cancel_date == date(2026, 10, 5)

    def test_short_row_does_not_fail(self):
        def build(doc, table):
            row = table.add_row()
            row.cells[2].text, row.cells[3].text = "02.10.26", "Имитрон 15 мин"
            # Удаляем хвостовые ячейки строки — так бывает в документах, правленных руками.
            for tc in row._tr.tc_lst[4:]:
                row._tr.remove(tc)

        sheet = parse_sheet(document(build))

        assert physio(sheet) == ["Имитрон 15 мин"]
        assert sheet.rows[0].cancel_date is None

    def test_empty_table_before_the_real_one(self):
        def build(doc, table):
            row = table.add_row()
            row.cells[2].text, row.cells[3].text = "02.10.26", "Имитрон"
            empty = doc.add_table(rows=0, cols=2)
            table._tbl.addprevious(empty._tbl)

        assert physio(parse_sheet(document(build))) == ["Имитрон"]

    def test_title_row_above_header(self):
        def build(doc, table):
            row = table.add_row()
            row.cells[2].text, row.cells[3].text = "02.10.26", "Имитрон"
            title = table.add_row()
            title._tr.getparent().remove(title._tr)
            table._tbl.tr_lst[0].addprevious(title._tr)
            title.cells[0].text = "Назначения"

        assert physio(parse_sheet(document(build))) == ["Имитрон"]

    def test_absurd_dosage_is_a_warning(self):
        sheet = parse_sheet(make_sheet(rows=[("02.10.26", "Имитрон 40000 мин 99 р/д", "")]))

        assert (sheet.rows[0].duration_min, sheet.rows[0].per_day) == (None, 1)
        assert len(sheet.warnings) == 2

    def test_course_starts_by_physio_not_by_consultation(self):
        sheet = parse_sheet(
            make_sheet(rows=[("05.10.26", "Имитрон", "")], diagnostics=["Консультация терапевта"])
        )

        assert sheet.start_date == date(2026, 10, 5)


# --- Листы разных врачей (TZ.md, решение 71) ---------------------------------------------------


def one_row(text: str):
    (row,) = parse_sheet(make_sheet(rows=[("05.10.26", text, "")])).rows
    return row


class TestDoctorVariants:
    @pytest.mark.parametrize(("room", "expected"), [("Палата 1", "1"), ("Палата 5а", "5а")])
    def test_room_without_number_sign(self, room, expected):
        header = f"ФИО  Тестова Анна Сергеевна     ИБ №7      ШРМ 4      Диагноз: I69      {room}"

        assert parse_sheet(make_sheet(header=header)).room == expected

    @pytest.mark.parametrize(
        ("text", "per_day"),
        [
            ("Thera Trainer Tigo, по 15 минут 2 раза в день, е/д", 2),
            ("Индивидуальное занятие ЛФК, 30 минут, 1 раз в день", 1),
            ("Индивид.ЛФК 2р/день, ежедневно", 2),
        ],
    )
    def test_frequency_wording(self, text, per_day):
        assert one_row(text).per_day == per_day

    def test_number_without_frequency(self):
        row = one_row("Индивидуальное занятие ЛФК 30 мин №2")

        assert (row.per_day, row.number) == (1, 2), "что значит №2 — решает импорт по процедуре"

    @pytest.mark.parametrize(
        "text",
        ["Массаж н/к 15 мин 8:45 №       246 кабинет", "Имитрон №2 2 р/д"],
    )
    def test_cabinet_or_explicit_frequency_is_not_a_number(self, text):
        assert one_row(text).number is None

    def test_group_list_without_dashes_last_line_with_dosage(self):
        text = (
            "Групповые занятия ЛФК:\nЭрго общая\nСпина в 16:30 - 181 каб.\n"
            "Колено в 15:00 – 179 каб.\nI can нога – 4р/день, ежедневно, по будням"
        )

        rows = parse_sheet(make_sheet(rows=[("07.10.26", text, "")])).rows

        assert [(r.raw_text, r.per_day, r.at) for r in rows] == [
            ("Групповые занятия ЛФК: Эрго общая", 1, ""),
            ("Групповые занятия ЛФК: Спина в 16:30 - 181 каб.", 1, "16:30"),
            ("Групповые занятия ЛФК: Колено в 15:00 – 179 каб.", 1, "15:00"),
            ("Групповые занятия ЛФК: I can нога", 1, ""),
        ]

    def test_bulleted_list_keeps_dosage_line_as_separate_item(self):
        """В списке с дефисами строка «название + дозировка» — отдельное назначение."""
        text = "Групповые занятия ЛФК\n- Эрго общая\nИндивидуальное занятие ЛФК 30 мин 1 р/д"

        rows = parse_sheet(make_sheet(rows=[("05.10.26", text, "")])).rows

        assert [r.raw_text for r in rows] == [
            "Групповые занятия ЛФК: Эрго общая",
            "Индивидуальное занятие ЛФК 30 мин 1 р/д",
        ]

    @pytest.mark.parametrize(("text", "at"), [("Колено в 9:05", "09:05"), ("Колено 25:00", "")])
    def test_time(self, text, at):
        assert one_row(text).at == at


class TestDosageInGroupTitle:
    """«ЛФК в группе, 30 минут, 3 раза в день» и группы ниже (решение 78)."""

    TEXT = "ЛФК в группе, 30 минут, ежедневно, 3 раза в день\nЭрго общая\nI-can н/к\nЭрго в/к"

    def test_groups_get_title_dosage_once_a_day(self):
        sheet = parse_sheet(make_sheet(rows=[("07.10.26", self.TEXT, "")]))

        assert [(r.raw_text, r.match_text, r.duration_min, r.per_day) for r in sheet.rows] == [
            ("ЛФК в группе: Эрго общая", "Эрго общая", 30, 1),
            ("ЛФК в группе: I-can н/к", "I-can н/к", 30, 1),
            ("ЛФК в группе: Эрго в/к", "Эрго в/к", 30, 1),
        ]
        assert sheet.warnings == []

    def test_frequency_not_matching_groups_is_a_warning(self):
        text = "ЛФК в группе, 30 минут, 4 раза в день\nЭрго общая\nI can нога"

        sheet = parse_sheet(make_sheet(rows=[("07.10.26", text, "")]))

        assert [r.per_day for r in sheet.rows] == [1, 1]
        assert any("4 р/д на 2 группы" in w for w in sheet.warnings)

    def test_title_without_groups_stays_a_row(self):
        """Групп нет — строка остаётся, врач выберет процедуру сам."""
        text = "ЛФК в группе, 30 минут, 1 раз в день"

        rows = parse_sheet(make_sheet(rows=[("07.10.26", text, "")])).rows

        assert [(r.raw_text, r.duration_min) for r in rows] == [(text, 30)]

    def test_next_title_closes_list(self):
        text = (
            "ЛФК в группе, 30 минут, 1 раз в день\nЭрго общая\n"
            "Бассейн\n- верхняя конечность\n30 мин 1 р/д"
        )

        rows = parse_sheet(make_sheet(rows=[("07.10.26", text, "")])).rows

        assert [r.raw_text for r in rows] == [
            "ЛФК в группе: Эрго общая",
            "Бассейн: верхняя конечность",
        ]


@pytest.mark.parametrize("text", ["Школа пациента", "ФТП", "ФТП.", "Школа пациента 1 раз в день"])
def test_not_rehab_rows_are_skipped(text):
    rows = parse_sheet(make_sheet(rows=[("07.10.26", text, ""), ("07.10.26", "Имитрон", "")])).rows

    assert [r.raw_text for r in rows] == ["Имитрон"]


def test_physiotherapy_with_details_is_kept():
    """Пропускается только общая строка «ФТП», уточнённая — назначение."""
    assert one_row("ФТП: Алмаг 15 мин").raw_text == "ФТП: Алмаг 15 мин"
