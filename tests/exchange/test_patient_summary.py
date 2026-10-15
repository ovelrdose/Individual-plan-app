"""Разбор сводки пациентов поста (apps.exchange.patient_summary). Данные вымышленные, структура —
как у сводок постов ОМР № 4 (12 колонок, палата в первой)."""

from datetime import date
from io import BytesIO

import pytest
from docx import Document

from apps.exchange.patient_summary import INTENSIVE_CARE, SummaryError, read_summary

HEADER_POST_1 = [
    "+",
    "ФИО",
    "В",
    "М/Ж",
    "№\nИБ",
    "стол",
    "Диагноз",
    "Бюджет",
    "Дата поступ.",
    "Дата\nвыпис.",
    "К/Д",
    "Л/В",
]
HEADER_POST_2 = [
    "№\nп/п",
    "\nФИО",
    "В",
    "М/Ж",
    "№\nИБ",
    "Cтол",
    "Диагноз",
    "Бюджет",
    "Дата\nпоступ.",
    "Дата\nвыписки",
    "К/Д",
    "Лечащий\nврач",
]


def row(room: str, name: str, ib: str = "", admitted: str = "", discharge: str = "", **extra):
    return [
        room,
        name,
        extra.get("age", "50"),
        extra.get("sex", "М"),
        ib,
        "10",
        "Х00.0",
        "ОМС",
        admitted,
        discharge,
        "3",
        extra.get("doctor", "Докторова А.А."),
    ]


def make_summary(
    rows: list[list[str]], header: list[str] = HEADER_POST_1, title: str = ""
) -> BytesIO:
    document = Document()
    document.add_paragraph(title or "Сводка пациентов ОМР 4 пост 1 на 15.10.26")
    document.add_paragraph("В отделении: 3")
    table = document.add_table(rows=1, cols=len(header))
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = text
    for values in rows:
        for cell, text in zip(table.add_row().cells, values, strict=True):
            cell.text = text
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


class TestReadSummary:
    def test_patient_fields(self):
        summary = read_summary(
            make_summary(
                [row("1", "Тестов И.П.", "10001", "02.10.26", "16.10.26", age="66", sex="М ")]
            )
        )

        assert summary.day == date(2026, 10, 15)
        (patient,) = summary.patients
        assert patient.full_name == "Тестов И.П."
        assert patient.surname == "Тестов"
        assert (patient.room, patient.department) == ("1", "")
        assert (patient.sex, patient.age, patient.history_number) == ("М", 66, "10001")
        assert (patient.admitted, patient.discharge) == (date(2026, 10, 2), date(2026, 10, 16))
        assert patient.doctor == "Докторова А.А."

    def test_second_post_header(self):
        summary = read_summary(
            make_summary(
                [row("9", "Пробная Е.А.", "10002", "28.09.26", "18.10.26", sex="Ж")],
                header=HEADER_POST_2,
            )
        )

        assert [p.full_name for p in summary.patients] == ["Пробная Е.А."]

    def test_names_normalized(self):
        summary = read_summary(
            make_summary(
                [
                    row("2", "Тестова К.А", "1", "02.10.26"),
                    row("2", "Пробнов В. В.", "2", "02.10.26"),
                    row("2", "Образцов И.С. (СВО)", "3", "02.10.26"),
                    row("8п ", "Примеров С.А. +сопр", "4", "12.10.26"),
                    row("9", "Безинициалова ", "5", "14.10.26", doctor="Врачова А,А,"),
                ]
            )
        )

        assert [p.full_name for p in summary.patients] == [
            "Тестова К.А.",
            "Пробнов В.В.",
            "Образцов И.С.",
            "Примеров С.А.",
            "Безинициалова",
        ]
        assert summary.patients[-1].doctor == "Врачова А.А."

    def test_rooms(self):
        summary = read_summary(
            make_summary(
                [
                    row("5а", "Первый А.А.", "1", "02.10.26"),
                    row(" / 3", "Второй Б.Б.", "2", "02.10.26"),
                    row("8п ", "Третий В.В.", "3", "02.10.26"),
                    row("П\nИ\nТ", "Четвёртый Г.Г.", "4", "02.10.26"),
                    row("8п(омр2)", "Пятая Д.Д.", "5", "02.10.26"),
                    row("Омр3", "Шестой Е.Е (6 п)", "6", "02.10.26"),
                    row("Омр3", "Седьмой Ж. Ж. ( 9 п)", "7", "02.10.26"),
                ]
            )
        )

        assert [(p.room, p.department) for p in summary.patients] == [
            ("5а", ""),
            ("3", ""),
            ("8", ""),
            (INTENSIVE_CARE, ""),
            ("8", "ОМР 2"),
            ("6", "ОМР 3"),
            ("9", "ОМР 3"),
        ]
        assert summary.patients[5].full_name == "Шестой Е.Е."

    def test_rows_without_admission_skipped(self):
        summary = read_summary(
            make_summary(
                [
                    row("5", "Первый А.А.", "1", "02.10.26"),
                    row("5", ""),
                    row("5а", "+сопр"),
                    row("7", "Второй Б.Б. (омр3)"),
                    row("8п", "Третий В.В. +сопр", "3", "12.10.26"),
                    row("8п", "Третий В.В. +сопр"),  # повтор объединённой ячейки
                    row("8п", "Третий В.В. +сопр", "3", "12.10.26"),
                ]
            )
        )

        assert [p.full_name for p in summary.patients] == ["Первый А.А.", "Третий В.В."]
        assert summary.skipped == ["Второй Б.Б."]

    def test_dates_and_numbers(self):
        summary = read_summary(
            make_summary(
                [
                    row("6", "Первый А.А.", "1", "06.10.26.", "", age="", sex="?"),
                    row("6", "Второй Б.Б.", "2", "06.10.2026", "31.02.26"),
                ]
            )
        )

        first, second = summary.patients
        assert first.admitted == date(2026, 10, 6) and first.discharge is None
        assert (first.age, first.sex) == (None, "")
        assert second.admitted == date(2026, 10, 6) and second.discharge is None

    def test_short_rows(self):
        cells = row("1", "Первый А.А.", "1", "02.10.26")
        document = Document()
        table = document.add_table(rows=2, cols=len(HEADER_POST_1))
        for cell, text in zip(table.rows[0].cells, HEADER_POST_1, strict=True):
            cell.text = text
        for cell, text in zip(table.rows[1].cells, cells, strict=True):
            cell.text = text
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)

        summary = read_summary(buffer)

        assert summary.day is None
        assert [p.full_name for p in summary.patients] == ["Первый А.А."]

    def test_no_patient_table(self):
        document = Document()
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Что-то"
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)

        with pytest.raises(SummaryError, match="ФИО"):
            read_summary(buffer)

    def test_not_docx(self):
        with pytest.raises(SummaryError, match="не читается"):
            read_summary(BytesIO(b"not a docx"))
