"""Варианты реальных листов назначений и сценарии из ревью (TZ.md §6)."""

from datetime import date
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from docx import Document

from apps.exchange.prescription_sheet import SheetError, parse_sheet
from apps.exchange.services import find_duplicates, import_prescription_sheet
from apps.exchange.views import PENDING_KEY
from apps.programs.models import Program
from tests.sheets import COLUMNS, HEADER, make_sheet

HEADER_TEXT = HEADER.format(fio="Тестова Анна Сергеевна", ib="1", shrm=4, dx="I69", room="5")


def document(build) -> BytesIO:
    """Лист с таблицей, которую собирает build(table) — для вариантов, которых нет в make_sheet."""
    doc = Document()
    doc.add_paragraph(HEADER_TEXT)
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    for cell, title in zip(table.rows[0].cells, COLUMNS, strict=True):
        cell.text = title
    build(doc, table)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def physio(sheet):
    return [r.raw_text for r in sheet.rows if not r.consultation]


class TestTableVariants:
    def test_vertically_merged_physio_cell_is_read_once(self):
        def build(doc, table):
            first, second = table.add_row().cells, table.add_row().cells
            first[2].text, first[3].text = "02.10.26", "Имитрон 15 мин 1 р/д"
            first[3].merge(second[3])

        assert physio(parse_sheet(document(build))) == ["Имитрон 15 мин 1 р/д"]

    def test_short_row_does_not_fail(self):
        def build(doc, table):
            row = table.add_row()
            row.cells[2].text, row.cells[3].text = "02.10.26", "Имитрон 15 мин"
            # Удаляем хвостовые ячейки строки — так бывает в документах, правленных руками.
            for tc in row._tr.tc_lst[4:]:
                row._tr.remove(tc)

        sheet = parse_sheet(document(build))

        assert physio(sheet) == ["Имитрон 15 мин"]
        assert sheet.rows[0].cancel_date is None

    def test_empty_table_before_the_real_one(self):
        def build(doc, table):
            row = table.add_row()
            row.cells[2].text, row.cells[3].text = "02.10.26", "Имитрон"
            empty = doc.add_table(rows=0, cols=2)
            table._tbl.addprevious(empty._tbl)

        assert physio(parse_sheet(document(build))) == ["Имитрон"]

    def test_title_row_above_header(self):
        def build(doc, table):
            row = table.add_row()
            row.cells[2].text, row.cells[3].text = "02.10.26", "Имитрон"
            title = table.add_row()
            title._tr.getparent().remove(title._tr)
            table._tbl.tr_lst[0].addprevious(title._tr)
            title.cells[0].text = "Назначения"

        assert physio(parse_sheet(document(build))) == ["Имитрон"]

    def test_consultation_list_continues_on_next_line(self):
        sheet = parse_sheet(
            make_sheet(
                diagnostics=[
                    "Консультация терапевта, медицинского психолога,\nлогопеда, эрготерапевта\n"
                    "Гликемический профиль"
                ]
            )
        )

        assert [r.match_text for r in sheet.rows if r.consultation] == [
            "терапевта",
            "медицинского психолога",
            "логопеда",
            "эрготерапевта",
        ]

    def test_extra_line_after_group_list(self):
        sheet = parse_sheet(
            make_sheet(
                rows=[
                    (
                        "02.10.26",
                        "Групповые занятия ЛФК\n- Эрго общая\n30 мин 2 р/д\nМассаж спины",
                        "",
                    )
                ]
            )
        )

        assert physio(sheet) == ["Групповые занятия ЛФК: Эрго общая", "Массаж спины"]
        assert (sheet.rows[0].duration_min, sheet.rows[0].per_day) == (30, 1)
        assert (sheet.rows[1].duration_min, sheet.rows[1].per_day) == (None, 1)
        assert "2 р/д на 1 группу" in sheet.warnings[0]

    def test_absurd_dosage_is_a_warning(self):
        sheet = parse_sheet(make_sheet(rows=[("02.10.26", "Имитрон 40000 мин 99 р/д", "")]))

        assert (sheet.rows[0].duration_min, sheet.rows[0].per_day) == (None, 1)
        assert len(sheet.warnings) == 2

    def test_course_starts_by_physio_not_by_consultation(self):
        sheet = parse_sheet(
            make_sheet(rows=[("05.10.26", "Имитрон", "")], diagnostics=["Консультация терапевта"])
        )

        assert sheet.start_date == date(2026, 10, 5)


class TestHeaderVariants:
    @pytest.mark.parametrize(
        ("room", "expected"),
        [("5а", "5а"), ("5 а", "5а"), ("5А", "5а"), ("9.", "9"), ("12", "12"), ("ВИП", "ВИП")],
    )
    def test_room(self, room, expected):
        assert parse_sheet(make_sheet(room=room)).room == expected

    def test_room_followed_by_word(self):
        header = "ФИО  Тестова Анна Сергеевна      ШРМ 4      Палата № 5 Диагноз: I69"

        assert parse_sheet(make_sheet(header=header)).room == "5"

    @pytest.mark.parametrize("kwargs", [{"room": "12345678901234"}, {"ib": "1" * 25}])
    def test_too_long_values(self, kwargs):
        with pytest.raises(SheetError, match="слишком длинные"):
            parse_sheet(make_sheet(**kwargs))

    def test_tabs_in_header(self):
        header = "ФИО\tТестова Анна Сергеевна\t\tИБ №7\tШРМ 3\tПалата № 2"

        sheet = parse_sheet(make_sheet(header=header))

        assert (sheet.full_name, sheet.history_number, sheet.shrm, sheet.room) == (
            "Тестова Анна Сергеевна",
            "7",
            3,
            "2",
        )


@pytest.mark.usefixtures("full_catalog")
class TestDuplicates:
    def sheet(self, **kwargs):
        return parse_sheet(make_sheet(**kwargs))

    def test_manual_program_without_ib_matches_by_name(self, make_program, department):
        make_program(
            full_name="Тестова Анна Сергеевна", history_number="", start_date=date(2026, 10, 2)
        )

        found = find_duplicates(
            department,
            self.sheet(fio="ТЕСТОВА Анна  Сергеевна".replace("  ", " ")),
            date(2026, 10, 2),
        )

        assert len(found) == 1

    def test_yo_in_name(self, make_program, department):
        make_program(
            full_name="Семёнова Анна Сергеевна", history_number="", start_date=date(2026, 10, 2)
        )

        assert find_duplicates(
            department, self.sheet(fio="Семенова Анна Сергеевна"), date(2026, 10, 2)
        )

    def test_different_ib_is_another_course(self, make_program, department):
        make_program(
            full_name="Тестова Анна Сергеевна", history_number="999", start_date=date(2026, 10, 2)
        )

        assert find_duplicates(department, self.sheet(ib="5402"), date(2026, 10, 2)) == []

    def test_adjacent_courses_overlap_on_one_day(self, make_program, department):
        make_program(history_number="5402", shrm=4, start_date=date(2026, 9, 18))  # до 02.10

        assert find_duplicates(department, self.sheet(ib="5402"), date(2026, 10, 2))
        assert find_duplicates(department, self.sheet(ib="5402"), date(2026, 10, 3)) == []

    def test_absurd_dosage_import(self, doctor, department):
        program = import_prescription_sheet(
            doctor,
            department,
            self.sheet(rows=[("02.10.26", "Имитрон 40000 мин", "")]),
            attending_doctor=doctor,
        )

        assert program.prescriptions.get().duration_min == 15, "длительность из справочника"
        assert "не похожа на правду" in program.import_warnings[0]


@pytest.mark.usefixtures("full_catalog")
class TestPendingSheets:
    url = reverse("exchange:sheet_import")

    def upload(self, client, fio, ib):
        file = SimpleUploadedFile("лист.docx", make_sheet(fio=fio, ib=ib).getvalue())
        return client.post(self.url, {"file": file})

    def test_two_tabs_do_not_mix(self, client, doctor):
        client.force_login(doctor)
        self.upload(client, "Альфа Анна Сергеевна", "1")
        self.upload(client, "Бета Анна Сергеевна", "2")
        tab_a = self.upload(client, "Альфа Анна Сергеевна", "1")
        tab_b = self.upload(client, "Бета Анна Сергеевна", "2")

        client.post(reverse("exchange:sheet_import_force"), {"token": tab_a.context["token"]})

        created = Program.objects.latest("pk")
        assert created.full_name == "Альфа Анна Сергеевна"
        assert tab_b.context["token"] in client.session[PENDING_KEY]

    @pytest.mark.parametrize(
        "pending",
        [
            {"department": 1, "doctor": 1, "sheet": {"unknown": 1}},
            {"department": 1, "doctor": 10**9, "sheet": {}},
            "мусор",
        ],
    )
    def test_stale_or_broken_session(self, client, doctor, department, pending):
        client.force_login(doctor)
        session = client.session
        session[PENDING_KEY] = {"tok": pending}
        session.save()

        response = client.post(
            reverse("exchange:sheet_import_force"), {"token": "tok"}, follow=True
        )

        assert "Загрузка устарела" in response.content.decode()

    def test_long_header_in_upload(self, client, doctor):
        client.force_login(doctor)
        file = SimpleUploadedFile("лист.docx", make_sheet(room="12345678901234").getvalue())

        response = client.post(self.url, {"file": file})

        assert response.status_code == 200
        assert "слишком длинные" in response.content.decode()
