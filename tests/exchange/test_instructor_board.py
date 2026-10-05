"""Инструкторы и распорядок из шахматки (apps.exchange.instructor_board) и команда
load_board_instructors. Файл генерируется в тесте: фамилии и пациенты вымышленные."""

from datetime import date, time
from io import StringIO
from pathlib import Path

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import Workbook

from apps.catalog.models import Procedure
from apps.exchange.instructor_board import _slot, normalize_group, read_board_staff
from apps.programs import services as programs
from apps.programs.models import Prescription
from apps.staff.models import Instructor, InstructorDuty, ShiftPattern

SLOTS = ["9:10-9:40", "9:50-10:20", "10:30-11:00", "13:00-13:30", "15:00-15:30", "18", "18.4"]
OLD = ["Старов/ Прежний", "Ушедший"]
NOW = ["Альфа/ Бета", "Гамма", "Дельта"]
GROUPS = {"эрго общая": "Эрго общая", "i can нога": "I can нога", "i can рука": "I can рука"}


def day_sheet(wb: Workbook, title: str, header: list[str], cells: dict[tuple[int, int], str]):
    ws = wb.create_sheet(title)
    ws.append([title.replace(",", "."), *header, "st", "Имитрон"])
    for row, slot in enumerate(SLOTS):
        ws.append([slot] + [cells.get((row, col), "") for col in range(len(header))])


@pytest.fixture
def board_file(tmp_path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    day_sheet(wb, "01,09", OLD, {(0, 0): "Метод. Работа"})
    for day in range(2, 6):
        cells = {
            (0, 0): f"{day}п Пациентов инд",
            (2, 0): "i-can нога" if day < 5 else "",  # 3 из 4 дней
            (4, 0): "I can рука" if day < 4 else "",  # 2 из 4 — у Гаммы реже
            (4, 1): "I can рука" if day == 4 else "",
            (5, 0): "Метод. Работа",
            (6, 0): "Метод. Работа",
            (0, 1): "Эрго общая",
            (3, 1): "БОС",
            (4, 2): "Метод. Работа" if day == 2 else "",  # один раз — не постоянно
            (1, 2): "16п ОМР 3 Вымыслов",
            (3, 2): "БОС 13п Выдумкин",
        }
        day_sheet(wb, f"0{day},09", NOW, cells)
    wb.create_sheet("Лист3")
    path = tmp_path / "board.xlsx"
    wb.save(path)
    return path


def test_reads_current_staff_and_recurring_duties(board_file):
    staff = read_board_staff(board_file, GROUPS)

    assert staff.sheets == 4 and (staff.first_sheet, staff.last_sheet) == ("02,09", "05,09")
    assert [c.members for c in staff.columns] == [["Альфа", "Бета"], ["Гамма"], ["Дельта"]]
    pair, gamma, delta = staff.columns
    assert [(d.slot, d.kind, d.group) for d in pair.duties] == [
        (time(10, 30), "GROUP_LEAD", "I can нога"),
        (time(15), "GROUP_LEAD", "I can рука"),
        (time(18), "METHOD_WORK", ""),
        (time(18, 40), "METHOD_WORK", ""),
    ]
    assert [(d.slot, d.kind, d.group) for d in gamma.duties] == [
        (time(9, 10), "GROUP_LEAD", "Эрго общая"),
        (time(13), "BOS", ""),
    ]
    assert delta.duties == [], "разовая «Метод. работа» и пациенты — не распорядок"


def test_slot_parsing():
    assert _slot("9:10-9:40") == time(9, 10)
    assert _slot("18") == time(18) and _slot("18.4") == time(18, 40)
    assert _slot(time(19, 20)) == time(19, 20)
    assert _slot("Итого") is None and _slot(None) is None and _slot("5") is None


def test_group_names_normalized():
    assert normalize_group("I-can нога") == normalize_group("i can  нога")
    assert normalize_group("Вест.гимнастика") == "вест гимнастика"


def test_empty_file_is_error(tmp_path):
    wb = Workbook()
    path = tmp_path / "empty.xlsx"
    wb.save(path)

    with pytest.raises(ValueError, match="нет листов"):
        read_board_staff(path, GROUPS)


@pytest.mark.usefixtures("full_catalog")
class TestCommand:
    def run(self, path, *args) -> str:
        out = StringIO()
        call_command(
            "load_board_instructors", str(path), "--since", "2026-10-05", *args, stdout=out
        )
        return out.getvalue()

    def test_loads_pairs_shifts_and_duties_idempotently(self, board_file):
        self.run(board_file)
        self.run(board_file)

        names = Instructor.objects.order_by("display_order").values_list("short_name", flat=True)
        assert list(names) == ["Альфа", "Бета", "Гамма", "Дельта"]
        alpha, beta = (
            Instructor.objects.get(short_name="Альфа"),
            Instructor.objects.get(short_name="Бета"),
        )
        assert alpha.partner == beta and alpha.team_label == "Альфа/ Бета"
        assert {p.pattern for p in ShiftPattern.objects.filter(instructor__in=[alpha, beta])} == {
            "2/2"
        }
        assert ShiftPattern.objects.get(instructor=beta).anchor_date == date(2026, 10, 7)
        assert ShiftPattern.objects.get(instructor__short_name="Гамма").pattern == "5/2"
        assert InstructorDuty.objects.filter(instructor=alpha).count() == 4
        assert InstructorDuty.objects.filter(instructor=beta).count() == 4
        gamma = InstructorDuty.objects.filter(instructor__short_name="Гамма")
        assert {(d.slot.start, d.kind) for d in gamma} == {
            (time(9, 10), "GROUP_LEAD"),
            (time(13), "BOS"),
        }

    def test_replace_removes_others_and_replans(
        self, board_file, settings, doctor, rehab, admin_user, make_program
    ):
        settings.DEBUG = True
        call_command("seed_dev", stdout=StringIO())
        program = make_program()
        programs.add_prescription(
            doctor,
            Prescription(
                program=program,
                procedure=Procedure.objects.get(name="Индивидуальное занятие", department=None),
            ),
        )
        assert program.bookings.filter(instructor__short_name="Соколов").exists()

        out = self.run(board_file, "--replace")

        assert "Удалены:" in out and "Соколов" in out
        assert set(Instructor.objects.values_list("short_name", flat=True)) == {
            "Альфа", "Бета", "Гамма", "Дельта",
        }  # fmt: skip
        names = set(
            program.bookings.filter(instructor__isnull=False).values_list(
                "instructor__short_name", flat=True
            )
        )
        assert names and names <= {"Альфа", "Бета", "Гамма", "Дельта"}
        # seed_dev больше не добавляет вымышленных к настоящим.
        call_command("seed_dev", stdout=StringIO())
        assert Instructor.objects.count() == 4

    def test_missing_file(self, tmp_path):
        with pytest.raises(CommandError, match="Нет файла"):
            self.run(tmp_path / "нет.xlsx")
