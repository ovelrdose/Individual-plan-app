"""QA шага 1.3: импорт листа назначений (TZ.md §5, §6) и выгрузка карты после импорта (§8).

Варианты реальных листов собираются из обезличенного make_sheet и правятся через python-docx.
"""

from datetime import date, datetime
from datetime import time as dt_time
from io import BytesIO
from urllib.parse import unquote

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from docx import Document
from openpyxl import load_workbook

from apps.accounts.access import SESSION_DEPARTMENT_KEY
from apps.accounts.models import Role
from apps.cards.models import CardExport
from apps.exchange.prescription_sheet import parse_sheet
from apps.exchange.services import DuplicateProgramError, import_prescription_sheet
from apps.exchange.views import PENDING_KEY
from apps.programs.models import Program
from tests.sheets import make_sheet

IMPORT_URL = reverse("exchange:sheet_import")
FORCE_URL = reverse("exchange:sheet_import_force")
PHYSIO, CANCEL = 3, 4  # колонки make_sheet


def header(fio: str = "Тестова Анна Сергеевна", rest: str = "ИБ №5402  ШРМ 4  Палата № 5а") -> str:
    return f"ФИО  {fio}      {rest}"


def to_upload(buffer: BytesIO, name: str = "лист.docx") -> SimpleUploadedFile:
    return SimpleUploadedFile(name, buffer.getvalue())


def edited(buffer: BytesIO, change) -> BytesIO:
    """Лист make_sheet, изменённый функцией change(document) — объединения, обрезанные строки."""
    document = Document(buffer)
    change(document)
    result = BytesIO()
    document.save(result)
    result.seek(0)
    return result


def do_import(user, department, buffer: BytesIO, *, force: bool = False) -> Program:
    return import_prescription_sheet(
        user, department, parse_sheet(buffer), attending_doctor=user, force=force
    )


def labels(program: Program) -> list[str]:
    return [p.label for p in program.prescriptions.all()]


def select(client, department) -> None:
    session = client.session
    session[SESSION_DEPARTMENT_KEY] = department.pk
    session.save()


# --- Шапка (FR-IMP-1) -------------------------------------------------------------------


class TestHeaderVariants:
    @pytest.mark.parametrize(
        "text",
        [
            "ФИО\tТестова Анна Сергеевна\tИБ №5402\tШРМ 4\tДиагноз: Т91.3\tПалата № 5а",
            "ФИО\xa0\xa0Тестова Анна Сергеевна\xa0\xa0ИБ\xa0№5402\xa0ШРМ\xa04\xa0"
            "Диагноз:\xa0Т91.3\xa0Палата\xa0№\xa05а",
            "ФИО Тестова Анна Сергеевна ИБ № 5402 ШРМ 4 Диагноз: Т91.3 Палата № 5а",
        ],
        ids=["tabs", "nbsp", "single-spaces"],
    )
    def test_separators(self, text):
        sheet = parse_sheet(make_sheet(header=text))

        assert (sheet.full_name, sheet.history_number, sheet.shrm, sheet.room) == (
            "Тестова Анна Сергеевна",
            "5402",
            4,
            "5а",
        )
        assert sheet.warnings == []

    def test_fields_in_other_order(self):
        sheet = parse_sheet(
            make_sheet(header="ФИО  Тестов Тест Тестович  Палата № 12  ИБ № 77  ШРМ 5")
        )

        assert (sheet.full_name, sheet.history_number, sheet.shrm, sheet.room) == (
            "Тестов Тест Тестович",
            "77",
            5,
            "12",
        )

    def test_double_surname(self):
        sheet = parse_sheet(make_sheet(fio="Тестова-Пробная Анна Сергеевна"))

        assert sheet.full_name == "Тестова-Пробная Анна Сергеевна"

    @pytest.mark.parametrize("bad", ["ШРМ 2", "ШРМ: 4", "ШРМ"])
    def test_bad_shrm_is_a_clear_error(self, client, doctor, bad):
        client.force_login(doctor)
        text = header(rest=f"ИБ №5402  {bad}  Палата № 5а")

        response = client.post(IMPORT_URL, {"file": to_upload(make_sheet(header=text))})

        assert response.status_code == 200
        assert "не найдено: ШРМ" in response.content.decode()

    def test_double_space_inside_full_name(self):
        sheet = parse_sheet(make_sheet(header=header(fio="Тестова  Анна Сергеевна")))

        assert sheet.full_name.split() == ["Тестова", "Анна", "Сергеевна"]

    def test_room_with_space_before_letter(self):
        sheet = parse_sheet(make_sheet(room="5 а"))

        assert sheet.room.replace(" ", "") == "5а"


# --- Таблица (FR-IMP-2…5) ---------------------------------------------------------------


class TestTableVariants:
    @pytest.mark.parametrize(
        ("prescribed", "cancel"),
        [("2.10.26", "5.10.26"), ("02.10.2026", "05.10.2026"), ("2.10.2026", "05.10.26")],
    )
    def test_date_formats(self, prescribed, cancel):
        (row,) = parse_sheet(make_sheet(rows=[(prescribed, "Имитрон", cancel)])).rows

        assert (row.prescribed_on, row.cancel_date) == (date(2026, 10, 2), date(2026, 10, 5))

    def test_multi_paragraph_cell_with_blank_lines_and_nbsp(self):
        text = "\nИмитрон\xa0\xa015\xa0мин\n\n \n1\xa0р/д е/д\n"

        (row,) = parse_sheet(make_sheet(rows=[("02.10.26", text, "")])).rows

        assert (row.raw_text.split(), row.duration_min, row.per_day) == (
            ["Имитрон", "15", "мин", "1", "р/д", "е/д"],
            15,
            1,
        )

    def test_bullets_with_tab_and_em_dash(self):
        text = "Групповые занятия ЛФК\n-\tЭрго общая\n—  Нейро - тренинг\n30 мин 1 р/д"

        rows = parse_sheet(make_sheet(rows=[("02.10.26", text, "")])).rows

        assert [r.match_text for r in rows] == ["Эрго общая", "Нейро - тренинг"]

    def test_consultations_with_and_and_dot(self):
        rows = parse_sheet(
            make_sheet(diagnostics=["Консультация медицинского психолога и логопеда."])
        ).rows

        assert [r.match_text for r in rows if r.consultation] == [
            "медицинского психолога",
            "логопеда",
        ]

    def test_vertically_merged_physio_cell(self):
        def merge(document):
            table = document.tables[0]
            table.cell(1, PHYSIO).merge(table.cell(2, PHYSIO))

        buffer = edited(
            make_sheet(rows=[("02.10.26", "Имитрон 15 мин", ""), ("03.10.26", "", "")]), merge
        )

        assert [r.raw_text for r in parse_sheet(buffer).rows] == ["Имитрон 15 мин"]

    def test_horizontally_merged_header_does_not_crash(self):
        def merge(document):
            table = document.tables[0]
            table.cell(0, 0).merge(table.cell(0, 1))

        buffer = edited(make_sheet(rows=[("02.10.26", "Имитрон", "05.10.26")]), merge)

        rows = parse_sheet(buffer).rows
        assert rows[0].raw_text == "Имитрон"
        assert rows[0].cancel_date == date(2026, 10, 5)

    def test_short_table_row_is_not_500(self, client, doctor, full_catalog):
        def cut(document):
            tr = document.tables[0].rows[2]._tr
            for tc in tr.tc_lst[2:]:
                tr.remove(tc)

        buffer = edited(
            make_sheet(rows=[("02.10.26", "Имитрон", ""), ("02.10.26", "Алмаг", "")]), cut
        )
        client.force_login(doctor)

        response = client.post(IMPORT_URL, {"file": to_upload(buffer)})

        assert response.status_code in (200, 302)

    def test_groups_without_dashes(self, doctor, department, full_catalog):
        text = "Групповые занятия ЛФК\nЭрго общая\nI can нога\n30 мин 2 р/д е/д"

        program = do_import(doctor, department, make_sheet(rows=[("02.10.26", text, "")]))

        assert labels(program) == ["Эрго общая", "I can нога"]

    def test_groups_and_individual_in_one_cell(self, doctor, department, full_catalog):
        text = (
            "Групповые занятия ЛФК\n- Эрго общая\n30 мин 1 р/д\n"
            "Индивидуальное занятие ЛФК 30 мин 1 р/д"
        )

        program = do_import(doctor, department, make_sheet(rows=[("02.10.26", text, "")]))

        assert labels(program) == ["Эрго общая", "Индивидуальное занятие"]

    def test_two_group_blocks_in_one_cell(self):
        text = (
            "Групповые занятия ЛФК\n- Эрго общая\n30 мин 2 р/д\n"
            "Групповые занятия ЛФК\n- Нейро-тренинг\n15 мин 1 р/д"
        )

        rows = parse_sheet(make_sheet(rows=[("02.10.26", text, "")])).rows

        assert [(r.match_text, r.duration_min, r.per_day) for r in rows] == [
            ("Эрго общая", 30, 1),
            ("Нейро-тренинг", 15, 1),
        ]

    def test_pool_with_group_as_bullet(self, doctor, department, full_catalog):
        text = "Бассейн\n- верхняя конечность\n30 мин 1 р/д"

        program = do_import(doctor, department, make_sheet(rows=[("02.10.26", text, "")]))

        assert labels(program) == ["ЛФК в воде: верхняя конечность"]

    def test_consultations_with_colon(self):
        rows = parse_sheet(make_sheet(diagnostics=["Консультации: психолога, логопеда"])).rows

        assert [r.match_text for r in rows if r.consultation] == ["психолога", "логопеда"]

    def test_consultations_in_physio_column(self, doctor, department, full_catalog):
        buffer = make_sheet(rows=[("02.10.26", "Консультация психолога, логопеда", "")])

        program = do_import(doctor, department, buffer)

        assert labels(program) == ["Психолог", "Логопед"]


# --- Даты назначений (FR-IMP-6, FR-PRG-3) -----------------------------------------------


@pytest.mark.usefixtures("full_catalog")
class TestDatesOnImport:
    def test_cancel_before_prescribed(self, doctor, department):
        program = do_import(
            doctor,
            department,
            make_sheet(rows=[("02.10.26", "Имитрон", ""), ("06.10.26", "Степпер", "04.10.26")]),
        )
        stepper = program.prescriptions.get(raw_text="Степпер")

        assert (stepper.start_date, stepper.cancel_date) == (date(2026, 10, 6), None)
        assert any("«Степпер»: дата отмены" in w for w in program.import_warnings)

    def test_cancel_on_first_day(self, doctor, department):
        program = do_import(
            doctor, department, make_sheet(rows=[("02.10.26", "Имитрон", "02.10.26")])
        )

        assert program.prescriptions.get().cancel_date is None
        assert any("дата отмены" in w for w in program.import_warnings)

    def test_cancel_on_last_day_is_kept(self, doctor, department):
        # ШРМ 4 с 02.10: курс до 16.10 включительно, отмена 16.10 допустима.
        program = do_import(
            doctor, department, make_sheet(rows=[("02.10.26", "Имитрон", "16.10.26")])
        )

        assert program.prescriptions.get().cancel_date == date(2026, 10, 16)
        assert program.import_warnings == []

    def test_prescribed_after_course_end(self, doctor, department):
        program = do_import(
            doctor,
            department,
            make_sheet(shrm=3, rows=[("02.10.26", "Имитрон", ""), ("20.10.26", "Степпер", "")]),
        )
        stepper = program.prescriptions.get(raw_text="Степпер")

        assert program.end_date == date(2026, 10, 12)
        assert stepper.start_date is None
        assert any("дата назначения вне курса" in w for w in program.import_warnings)

    @pytest.mark.parametrize(
        ("shrm", "end"),
        [(3, date(2026, 10, 12)), (4, date(2026, 10, 16)), (5, date(2026, 10, 21))],
    )
    def test_course_length_by_shrm(self, doctor, department, shrm, end):
        program = do_import(doctor, department, make_sheet(shrm=shrm))

        assert (program.start_date, program.end_date) == (date(2026, 10, 2), end)
        assert len(program.course_dates()) == {3: 11, 4: 15, 5: 20}[shrm]


# --- Ввод, который не должен давать 500 (FR-IMP-12) -------------------------------------


@pytest.mark.usefixtures("full_catalog")
class TestNoServerError:
    def test_empty_file(self, client, doctor):
        client.force_login(doctor)

        response = client.post(IMPORT_URL, {"file": SimpleUploadedFile("лист.docx", b"")})

        assert response.status_code == 200
        assert response.context["form"].errors["file"]
        assert not Program.objects.exists()

    def test_upper_case_extension(self, client, doctor):
        client.force_login(doctor)

        response = client.post(IMPORT_URL, {"file": to_upload(make_sheet(), "ЛИСТ.DOCX")})

        assert response.status_code == 302
        assert Program.objects.count() == 1

    def test_huge_duration(self, client, doctor):
        client.force_login(doctor)
        buffer = make_sheet(rows=[("02.10.26", "Имитрон 40000 мин", "")])

        response = client.post(IMPORT_URL, {"file": to_upload(buffer)})

        assert response.status_code in (200, 302)

    @pytest.mark.parametrize(
        "text",
        [
            header(rest="ИБ №5402  ШРМ 4  Палата № 5а,6б,7в,8г"),
            header(fio="Тестова " * 25 + "Сергеевна"),
        ],
        ids=["long-room", "long-name"],
    )
    def test_too_long_values(self, client, doctor, text):
        client.force_login(doctor)

        response = client.post(IMPORT_URL, {"file": to_upload(make_sheet(header=text))})

        assert response.status_code in (200, 302)

    def test_only_consultations_without_physio(self, client, doctor):
        client.force_login(doctor)
        buffer = make_sheet(rows=[], diagnostics=["Консультация психолога"])

        response = client.post(IMPORT_URL, {"file": to_upload(buffer)})

        program = Program.objects.get()
        assert response.url == reverse("programs:detail", args=[program.pk])
        assert labels(program) == ["Психолог"]
        assert "начало курса взято сегодняшним днём" in program.import_warnings[0]


# --- Повтор (FR-IMP-11) -----------------------------------------------------------------


NO_IB = "ФИО  {fio}      ШРМ 3          Палата № 9"


@pytest.mark.usefixtures("full_catalog")
class TestDuplicates:
    def test_last_day_equals_first_day_is_duplicate(self, doctor, department):
        first = do_import(doctor, department, make_sheet(shrm=3))  # 02.10–12.10

        with pytest.raises(DuplicateProgramError) as error:
            do_import(doctor, department, make_sheet(rows=[("12.10.26", "Имитрон", "")]))
        assert error.value.programs == [first]

    def test_next_day_after_course_is_not_duplicate(self, doctor, department):
        do_import(doctor, department, make_sheet(shrm=3))  # 02.10–12.10

        do_import(doctor, department, make_sheet(rows=[("13.10.26", "Имитрон", "")]))

        assert Program.objects.count() == 2

    def test_new_course_ending_on_first_day_of_existing(self, doctor, department):
        do_import(doctor, department, make_sheet(rows=[("12.10.26", "Имитрон", "")]))

        with pytest.raises(DuplicateProgramError):
            # ШРМ 3 с 02.10 — до 12.10, ровно первый день существующего курса.
            do_import(doctor, department, make_sheet(shrm=3))

    def test_manually_extended_course_counts(self, doctor, department):
        from apps.programs.services import save_program

        first = do_import(doctor, department, make_sheet(shrm=3))
        first.end_date = date(2026, 10, 20)
        save_program(doctor, first, end_date_changed=True)

        with pytest.raises(DuplicateProgramError):
            do_import(doctor, department, make_sheet(rows=[("18.10.26", "Имитрон", "")]))

    def test_same_name_in_other_case_without_ib(self, doctor, department):
        do_import(doctor, department, make_sheet(header=NO_IB.format(fio="Тестов Тест Тестович")))

        with pytest.raises(DuplicateProgramError):
            do_import(
                doctor, department, make_sheet(header=NO_IB.format(fio="тестов тест тестович"))
            )

    def test_same_name_with_yo_without_ib(self, doctor, department):
        do_import(doctor, department, make_sheet(header=NO_IB.format(fio="Семёнов Пётр Иванович")))

        with pytest.raises(DuplicateProgramError):
            do_import(
                doctor, department, make_sheet(header=NO_IB.format(fio="Семенов Петр Иванович"))
            )

    def test_duplicate_page_links_existing(self, client, doctor):
        client.force_login(doctor)
        client.post(IMPORT_URL, {"file": to_upload(make_sheet())})
        existing = Program.objects.get()

        response = client.post(IMPORT_URL, {"file": to_upload(make_sheet())})

        assert reverse("programs:detail", args=[existing.pk]) in response.content.decode()
        assert Program.objects.count() == 1


# --- Отделение и сессия -----------------------------------------------------------------


@pytest.mark.usefixtures("full_catalog")
class TestDepartmentsAndSession:
    def test_doctor_with_two_departments_imports_into_selected(
        self, client, make_user, department, other_department
    ):
        user = make_user(
            "both", (department, Role.DOCTOR), (other_department, Role.DOCTOR), short_name="Д."
        )
        client.force_login(user)
        select(client, other_department)

        client.post(IMPORT_URL, {"file": to_upload(make_sheet())})

        program = Program.objects.get()
        assert (program.department, program.attending_doctor) == (other_department, user)

    def test_rehab_in_selected_department_cannot_import(
        self, client, make_user, department, other_department
    ):
        user = make_user("mixed", (department, Role.DOCTOR), (other_department, Role.REHAB))
        client.force_login(user)
        select(client, other_department)

        response = client.post(IMPORT_URL, {"file": to_upload(make_sheet())})

        # С 04.10 отделение выбирается явно: в «чужое по роли» отделение из шапки загрузить
        # нельзя, а молча подставить другое — тоже (форма просит выбрать отделение).
        assert response.status_code == 200
        assert "department" in response.context["form"].errors
        assert not Program.objects.exists()

    def test_admin_without_membership_uses_selected_department(
        self, client, admin_user, make_user, department, other_department
    ):
        other_doctor = make_user("od", (other_department, Role.DOCTOR))
        client.force_login(admin_user)
        select(client, other_department)

        client.post(
            IMPORT_URL, {"file": to_upload(make_sheet()), "attending_doctor": other_doctor.pk}
        )

        program = Program.objects.get()
        assert (program.department, program.attending_doctor) == (other_department, other_doctor)

    def test_admin_cannot_pick_doctor_of_other_department(
        self, client, admin_user, doctor, make_user, other_department
    ):
        other_doctor = make_user("od", (other_department, Role.DOCTOR))
        client.force_login(admin_user)
        select(client, doctor.memberships.get().department)

        response = client.post(
            IMPORT_URL, {"file": to_upload(make_sheet()), "attending_doctor": other_doctor.pk}
        )

        assert response.status_code == 200
        assert "attending_doctor" in response.context["form"].errors
        assert not Program.objects.exists()

    def _make_pending(self, client, doctor) -> str:
        client.force_login(doctor)
        client.post(IMPORT_URL, {"file": to_upload(make_sheet())})
        response = client.post(IMPORT_URL, {"file": to_upload(make_sheet())})
        assert PENDING_KEY in client.session
        return response.context["token"]

    def _tamper(self, client, token: str, **values) -> None:
        session = client.session
        session[PENDING_KEY][token].update(values)
        session.save()

    def test_force_twice_creates_one_program(self, client, doctor):
        token = self._make_pending(client, doctor)

        client.post(FORCE_URL, {"token": token})
        second = client.post(FORCE_URL, {"token": token}, follow=True)

        assert Program.objects.count() == 2
        assert "Загрузка устарела" in second.content.decode()

    def test_force_with_unknown_token(self, client, doctor):
        self._make_pending(client, doctor)

        response = client.post(FORCE_URL, {"token": "чужой"}, follow=True)

        assert "Загрузка устарела" in response.content.decode()
        assert Program.objects.count() == 1

    def test_force_after_switching_department_keeps_original(
        self, client, make_user, department, other_department
    ):
        user = make_user("both", (department, Role.DOCTOR), (other_department, Role.DOCTOR))
        select(client, department)
        token = self._make_pending(client, user)
        select(client, other_department)

        client.post(FORCE_URL, {"token": token})

        assert Program.objects.count() == 2
        assert set(Program.objects.values_list("department", flat=True)) == {department.pk}

    def test_tampered_department_of_stranger(self, client, doctor, other_department):
        token = self._make_pending(client, doctor)
        self._tamper(client, token, department=other_department.pk)

        response = client.post(FORCE_URL, {"token": token})

        assert response.status_code in (302, 404)
        assert Program.objects.count() == 1

    def test_tampered_department_where_user_is_rehab_is_403(
        self, client, make_user, department, other_department
    ):
        user = make_user("mixed", (department, Role.DOCTOR), (other_department, Role.REHAB))
        select(client, department)
        token = self._make_pending(client, user)
        self._tamper(client, token, department=other_department.pk)

        response = client.post(FORCE_URL, {"token": token})

        assert response.status_code == 403
        assert not Program.objects.filter(department=other_department).exists()

    def test_tampered_doctor_of_other_department(self, client, doctor, make_user, other_department):
        stranger = make_user("od", (other_department, Role.DOCTOR))
        token = self._make_pending(client, doctor)
        self._tamper(client, token, doctor=stranger.pk)

        response = client.post(FORCE_URL, {"token": token}, follow=True)

        assert "врачом этого отделения" in response.content.decode()
        assert Program.objects.count() == 1


# --- Карта после импорта (§8) -----------------------------------------------------------


FULL_ROWS = [
    ("02.10.26", "Групповые занятия ЛФК\n- Эрго общая\n- I can нога\n30 мин 2 р/д е/д", ""),
    ("02.10.26", "Индивидуальное занятие ЛФК 30 мин 1 р/д е/д", ""),
    ("02.10.26", "ST – 150 15 мин 1 р/д е/д", ""),
    ("04.10.26", "Степпер 15 минут 1 р/д, ежедневно", "06.10.26"),
]
CONSULTATIONS = "Консультация терапевта, медицинского психолога, эрготерапевта."
CARD_LABELS = ["Группа Эрго общая", "Группа I can нога", "Инд.занятие", "st-150", "Степпер"]
CELLS = {  # ФИО, пол/возраст, врач, отделение, строка дат, первая строка процедур
    3: ("I10", "I11", "J12", "N2", 17, 18),
    4: ("J10", "J11", "K12", "O2", 17, 18),
    5: ("J10", "J11", "K12", "O2", 18, 19),
}


def column(index: int) -> str:
    return chr(ord("C") + index)


@pytest.mark.usefixtures("full_catalog")
class TestCardAfterImport:
    @pytest.mark.parametrize("shrm", [3, 4, 5])
    def test_card_cells(self, client, doctor, shrm):
        client.force_login(doctor)
        buffer = make_sheet(shrm=shrm, rows=FULL_ROWS, diagnostics=[CONSULTATIONS])
        client.post(IMPORT_URL, {"file": to_upload(buffer)})
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]))

        assert response.status_code == 200
        sheet = load_workbook(BytesIO(response.content)).active
        name, sex_age, doc, dept, date_row, first_row = CELLS[shrm]
        assert sheet[name].value == "Тестова Анна Сергеевна"
        assert sheet[sex_age].value == "(Ж)"
        assert sheet[doc].value == "Иванова А.А."
        assert sheet[dept].value == "ОМР № 4"

        first_block = min(len(program.course_dates()), 15 if shrm == 4 else 11 if shrm == 3 else 10)
        dates = [sheet[f"{column(i)}{date_row}"] for i in range(first_block)]
        assert all(isinstance(cell.value, datetime) for cell in dates)
        assert all(cell.number_format == "dd.mm" for cell in dates)
        assert [cell.value.date() for cell in dates] == program.course_dates()[:first_block]

        expected = [*CARD_LABELS, "Психолог", "Эрготерапевт"]
        got = [sheet[f"A{first_row + i}"].value for i in range(len(expected) + 1)]
        assert got == [*expected, None]

        # Степпер с 04.10 по 05.10 (06.10 — первый день без процедуры).
        stepper_row = first_row + CARD_LABELS.index("Степпер")
        fills = [sheet[f"{column(i)}{stepper_row}"].fill.fill_type for i in range(6)]
        assert fills == ["solid", "solid", None, None, "solid", "solid"]
        assert all(sheet[f"{column(i)}{stepper_row}"].value is None for i in range(first_block))
        # Верхний блок «Расписание занятий» — типичный день: время настоящим значением Excel.
        top = [sheet[f"B{row}"].value for row in range(4, 14) if sheet[f"B{row}"].value is not None]
        assert top, "верхний блок заполнен расписанием (фаза 2)"
        assert all(isinstance(value, dt_time) for value in top)
        assert top == sorted(top)
        assert CardExport.objects.get().program == program

    def test_shrm5_second_block(self, client, doctor):
        client.force_login(doctor)
        client.post(IMPORT_URL, {"file": to_upload(make_sheet(shrm=5, rows=FULL_ROWS))})
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]))

        sheet = load_workbook(BytesIO(response.content)).active
        second = [sheet[f"{column(i)}31"].value for i in range(10)]
        assert all(isinstance(value, datetime) for value in second)
        assert [value.date() for value in second] == program.course_dates()[10:]
        assert [sheet[f"A{32 + i}"].value for i in range(len(CARD_LABELS))] == CARD_LABELS

    def test_unrecognized_after_import_blocks_card(self, client, doctor):
        client.force_login(doctor)
        buffer = make_sheet(rows=[("02.10.26", "Неизвестная процедура 10 мин", "")])
        client.post(IMPORT_URL, {"file": to_upload(buffer)})
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]), follow=True)

        assert response.redirect_chain[-1][0] == reverse("programs:detail", args=[program.pk])
        assert "Неизвестная процедура 10 мин" in response.content.decode()

    def test_too_many_rows_for_shrm3_is_a_message(self, client, doctor):
        client.force_login(doctor)
        rows = [("02.10.26", "Имитрон", "")] * 16
        client.post(IMPORT_URL, {"file": to_upload(make_sheet(shrm=3, rows=rows))})
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]), follow=True)

        assert "Процедур для карты 16" in response.content.decode()
        assert not CardExport.objects.exists()

    @pytest.mark.parametrize(
        ("fio", "surname"),
        [
            ("О'Тестова Анна Сергеевна", "О'Тестова"),
            ("Тестова-Пробная Анна Сергеевна", "Тестова-Пробная"),
        ],
    )
    def test_filename(self, client, doctor, fio, surname):
        client.force_login(doctor)
        client.post(IMPORT_URL, {"file": to_upload(make_sheet(fio=fio))})
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]))

        disposition = response["Content-Disposition"]
        encoded = disposition.split("filename*=UTF-8''", 1)[1]
        assert "'" not in encoded and " " not in encoded
        assert unquote(encoded) == f"Программа_{surname}_5402_02.10.2026.xlsx"
