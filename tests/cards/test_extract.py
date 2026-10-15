from pathlib import Path

import pytest
from django.conf import settings
from openpyxl import Workbook, load_workbook

from apps.cards.card_templates import bundled_mappings, bundled_template
from apps.cards.xlsx.extract import (
    column_width,
    data_cells,
    extract_templates,
    print_scale,
)

SOURCE = Path(settings.BASE_DIR) / "primary_docs" / "programm_template.xlsx"
MAPPINGS = bundled_mappings()


def assert_clean_template(path: Path, shrm: int) -> None:
    workbook = load_workbook(path)
    assert workbook.sheetnames == [f"ШРМ {shrm}"]
    sheet = workbook.active
    leftovers = {
        c: sheet[c].value for c in data_cells(MAPPINGS[shrm]) if sheet[c].value is not None
    }
    assert leftovers == {}, "в шаблоне остались примерные данные"
    assert len(sheet._images) == 1
    assert sheet.print_area.endswith(f"${MAPPINGS[shrm].date_blocks[-1].last_procedure_row}")
    assert_two_sided(sheet, shrm)


def assert_two_sided(sheet, shrm: int) -> None:
    """Лицевая сторона — расписание, шапка и «Уважаемый пациент!», оборотная — таблица дат."""
    front_last_row = MAPPINGS[shrm].date_blocks[0].date_row - 1
    assert [b.id for b in sheet.row_breaks.brk] == [front_last_row]
    # «Вписать в страницу» Excel применяет без ручных разрывов — масштаб числом.
    assert not sheet.sheet_properties.pageSetUpPr.fitToPage
    assert 30 <= sheet.page_setup.scale <= 100
    assert sheet.page_setup.orientation == "landscape"
    assert int(sheet.page_setup.paperSize) == int(sheet.PAPERSIZE_A4)
    front = [c.value for row in sheet.iter_rows(max_row=front_last_row) for c in row]
    assert "Уважаемый пациент!" in front
    # Перекрывающиеся описания колонок Excel читает непредсказуемо.
    covered: set[int] = set()
    for dim in sheet.column_dimensions.values():
        columns = set(range(dim.min, dim.max + 1))
        assert not columns & covered, f"колонка {dim.index} описана дважды"
        covered |= columns


@pytest.mark.skipif(not SOURCE.exists(), reason="нет исходного файла заказчика")
def test_extract_from_customer_file(tmp_path):
    paths = extract_templates(SOURCE, MAPPINGS, tmp_path)

    assert [p.name for p in paths] == ["shrm3.xlsx", "shrm4.xlsx", "shrm5.xlsx"]
    for shrm, path in zip((3, 4, 5), paths, strict=True):
        assert_clean_template(path, shrm)


@pytest.mark.parametrize("shrm", [3, 4, 5])
def test_bundled_templates_are_clean(shrm):
    assert_clean_template(bundled_template(shrm), shrm)


def test_column_width_reads_grouped_columns_without_creating_them():
    sheet = Workbook().active
    sheet.column_dimensions.group("C", "F", hidden=False)
    sheet.column_dimensions["C"].width = 7.5

    assert column_width(sheet, 5) == 7.5
    assert column_width(sheet, 10) == pytest.approx(8.43)
    assert list(sheet.column_dimensions) == ["C"]


def test_print_scale_fits_tallest_page():
    sheet = Workbook().active
    for row in range(1, 41):
        sheet.row_dimensions[row].height = 30
    sheet.page_margins.top = sheet.page_margins.bottom = 0.2

    one_page = print_scale(sheet, 10, [(1, 10)])
    two_pages = print_scale(sheet, 10, [(1, 10), (11, 40)])

    assert one_page == 100, "маленькая карта не увеличивается"
    # 30 строк по 30 пт = 900 пт на странице высотой 595 − 29 пт.
    assert two_pages == int((595 - 0.4 * 72) / 900 * 0.92 * 100)
