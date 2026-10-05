"""Выгрузка шахматки (TZ.md FR-CRD-7): файл читается обратно, проверяются значения и форматы."""

from datetime import date, time
from io import BytesIO

from openpyxl import load_workbook

from apps.cards.xlsx.board import BoardCell, BoardSheet, render_board

DAY = date(2026, 10, 5)


def sheet(**fields) -> BoardSheet:
    values = {
        "day": DAY,
        "slots": [(time(9, 10), time(9, 40)), (time(9, 50), time(10, 20))],
        "columns": ["Соколов", "Волков/ Лебедева"],
        "cells": [
            [BoardCell("9п Иванов"), BoardCell("Метод. работа", busy=True)],
            [BoardCell(), BoardCell("3п ОМР № 1 Петров")],
        ],
        "equipment": ["st-150", "Pablo"],
        "equipment_times": [time(13), time(13, 15), time(15, 30)],
        "equipment_cells": [
            [(["9п Иванов", "4п Сидоров"], False), ([], False)],
            [([], False), (["2п Орлов"], False)],
            [(["5п Лишний"], True), ([], False)],
        ],
    } | fields
    return BoardSheet(**values)


def read(data: BoardSheet):
    return load_workbook(BytesIO(render_board(data))).active


def test_layout_values_and_formats():
    ws = read(sheet())

    assert ws.title == "05.10"
    assert ws["A1"].value.date() == DAY and ws["A1"].number_format == "dd.mm.yyyy"
    assert [ws["B1"].value, ws["C1"].value] == ["Соколов", "Волков/ Лебедева"]
    assert [ws["A2"].value, ws["A3"].value] == ["9:10-9:40", "9:50-10:20"]
    assert ws["B2"].value == "9п Иванов" and ws["B3"].value is None
    assert ws["C2"].value == "Метод. работа" and ws["C2"].fill.start_color.rgb == "FFD9D9D9"
    assert ws["B2"].fill.fill_type is None
    assert ws["D1"].value is None, "пустая колонка перед тренажёрами"
    assert [ws["E1"].value, ws["F1"].value, ws["G1"].value] == ["Время", "st-150", "Pablo"]
    assert ws["E2"].value == time(13) and ws["E2"].number_format == "h:mm"
    assert ws["F2"].value == "9п Иванов\n4п Сидоров" and ws["F2"].alignment.wrap_text
    assert ws["G3"].value == "2п Орлов"
    assert ws["F4"].value == "5п Лишний" and ws["F4"].fill.start_color.rgb == "FFFFF2CC"
    assert ws.freeze_panes == "B2"
    assert ws.page_setup.orientation == "landscape"
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0


def test_without_equipment():
    ws = read(sheet(equipment=[], equipment_times=[], equipment_cells=[]))

    assert ws.max_column == 3
