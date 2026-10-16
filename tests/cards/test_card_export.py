from datetime import date, datetime
from datetime import time as dt_time
from io import BytesIO
from urllib.parse import unquote

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import Role
from apps.cards.models import CardExport
from apps.cards.services import build_card_data, card_filename, choose_template, export_card
from apps.cards.xlsx.formatting import sex_age_text
from apps.cards.xlsx.writer import CardRenderError
from apps.programs.models import Prescription, Program
from apps.programs.services import save_program
from tests.sheets import make_sheet


def upload(**kwargs):
    return SimpleUploadedFile("лист.docx", make_sheet(**kwargs).getvalue())


IMPORT_URL = reverse("exchange:sheet_import")
FULL_ROWS = [
    ("02.10.26", "Групповые занятия ЛФК\n- Эрго общая\n- I can нога\n30 мин 2 р/д е/д", ""),
    ("02.10.26", "Индивидуальное занятие ЛФК 30 мин 1 р/д е/д", ""),
    ("02.10.26", "ST – 150 15 мин 1 р/д е/д", ""),
    ("04.10.26", "Степпер 15 минут 1 р/д, ежедневно", "06.10.26"),
]
CONSULTATIONS = "Консультация терапевта, медицинского психолога, эрготерапевта."
CARD_LABELS = ["Группа Эрго общая", "Группа I can нога", "Инд.занятие", "st-150", "Степпер"]
CELLS = {  # ФИО, пол/возраст, врач, отделение, строка дат, первая строка процедур
    3: ("I10", "I11", "J12", "N2", 17, 18),
    4: ("J10", "J11", "K12", "O2", 17, 18),
    5: ("J10", "J11", "K12", "O2", 18, 19),
}


def column(index: int) -> str:
    return chr(ord("C") + index)


@pytest.fixture
def program(make_program, procedures):
    program = make_program(shrm=3, sex="Ж", age=17)  # 05.10–15.10, 11 дат
    Prescription.objects.create(program=program, procedure=procedures["group"], card_order=1)
    Prescription.objects.create(
        program=program,
        procedure=procedures["card_only"],
        card_order=2,
        start_date=date(2026, 10, 7),
        cancel_date=date(2026, 10, 10),
    )
    Prescription.objects.create(
        program=program, procedure=procedures["equipment"], card_order=3, in_card=False
    )
    return program


def test_card_data(program):
    data = build_card_data(program)

    assert (data.full_name, data.sex, data.age, data.doctor) == (
        "Тестов Тест Тестович",
        "Ж",
        17,
        "Иванова А.А.",
    )
    assert len(data.course_dates) == 11
    assert [(p.label, p.start_date, p.cancel_date) for p in data.procedures] == [
        ("Группа Эрго общая", None, None),
        ("Эрготерапевт", date(2026, 10, 7), date(2026, 10, 10)),
    ]
    assert data.schedule == ()


def test_unrecognized_rows_block_export(program):
    Prescription.objects.create(program=program, raw_text="Непонятная строка", card_order=4)

    with pytest.raises(CardRenderError, match="Непонятная строка"):
        build_card_data(program)


def test_unrecognized_row_outside_card_does_not_block(program):
    Prescription.objects.create(program=program, raw_text="Непонятная строка", in_card=False)

    build_card_data(program)


@pytest.mark.parametrize(
    ("sex", "age", "expected"),
    [("Ж", 17, "(Ж) 17 лет"), ("М", None, "(М)"), ("", 52, "52 года"), ("", None, "")],
)
def test_sex_age_text(sex, age, expected):
    assert sex_age_text(sex, age) == expected


def test_filename(program):
    assert card_filename(program) == "Программа_Тестов_100_05.10.2026.xlsx"
    program.history_number = ""
    assert card_filename(program) == "Программа_Тестов_05.10.2026.xlsx"


def test_download(client, doctor, program):
    client.force_login(doctor)

    response = client.get(reverse("cards:download", args=[program.pk]))

    assert response.status_code == 200
    assert "filename*=UTF-8''%D0%9F" in response["Content-Disposition"]
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet["I10"].value == "Тестов Тест Тестович"
    assert sheet["I11"].value == "(Ж) 17 лет, палата 9"
    assert isinstance(sheet["C17"].value, datetime)
    assert [sheet["A18"].value, sheet["A19"].value, sheet["A20"].value] == [
        "Группа Эрго общая",
        "Эрготерапевт",
        None,
    ]
    assert sheet["C19"].fill.fill_type == "solid", "Эрготерапевт ещё не начат 05.10"
    assert CardExport.objects.get().user == doctor


def test_download_error_returns_to_program(client, doctor, program):
    Prescription.objects.create(program=program, raw_text="Непонятная строка", card_order=4)
    client.force_login(doctor)

    response = client.get(reverse("cards:download", args=[program.pk]), follow=True)

    assert response.redirect_chain[-1][0] == reverse("programs:detail", args=[program.pk])
    assert "нераспознанных назначений" in response.content.decode()
    assert not CardExport.objects.exists()


def test_rehab_can_download_stranger_cannot(client, rehab, make_user, other_department, program):
    client.force_login(rehab)
    assert client.get(reverse("cards:download", args=[program.pk])).status_code == 200

    client.force_login(make_user("stranger", (other_department, Role.DOCTOR)))
    assert client.get(reverse("cards:download", args=[program.pk])).status_code == 404


def test_card_schedule_follows_slot_grid(program):
    from datetime import time

    from apps.catalog.models import InstructorSlot

    InstructorSlot.objects.create(start=time(9, 10), end=time(9, 40))
    InstructorSlot.objects.create(start=time(9, 50), end=time(10, 20))
    InstructorSlot.objects.create(start=time(18, 0), end=time(18, 30), is_evening=True)

    data = build_card_data(program)

    # Расписание ещё не подобрано — в карте только окна дневной сетки, вечерних слотов нет.
    assert [(item.start, item.label) for item in data.schedule] == [
        (time(9, 10), ""),
        (time(9, 50), ""),
    ]


def test_card_marks_admission_and_discharge_as_rest_dates(program):
    data = build_card_data(program)

    assert data.rest_dates == {date(2026, 10, 5), date(2026, 10, 15)}


@pytest.mark.usefixtures("full_catalog")
class TestCardAfterImport:
    """Лист назначений → программа → карта, как у врача (§8)."""

    @pytest.mark.parametrize("shrm", [3, 4, 5])
    def test_card_cells(self, client, doctor, shrm):
        client.force_login(doctor)
        client.post(
            IMPORT_URL,
            {"file": upload(shrm=shrm, rows=FULL_ROWS, diagnostics=[CONSULTATIONS])},
        )
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]))

        assert response.status_code == 200
        sheet = load_workbook(BytesIO(response.content)).active
        name, sex_age, doc, dept, date_row, first_row = CELLS[shrm]
        assert sheet[name].value == "Тестова Анна Сергеевна"
        assert sheet[sex_age].value == "(Ж), палата 5а"
        assert sheet[doc].value == "Иванова А.А."
        assert sheet[dept].value == "ОМР № 4"

        first_block = min(len(program.course_dates()), 15 if shrm == 4 else 11 if shrm == 3 else 10)
        dates = [sheet[f"{column(i)}{date_row}"] for i in range(first_block)]
        assert all(isinstance(cell.value, datetime) for cell in dates)
        assert all(cell.number_format == "dd.mm" for cell in dates)
        assert [cell.value.date() for cell in dates] == program.course_dates()[:first_block]

        expected = [*CARD_LABELS, "Психолог", "Эрготерапевт"]
        got = [sheet[f"A{first_row + i}"].value for i in range(len(expected) + 1)]
        assert got == [*expected, None]

        # Степпер с 04.10 по 05.10 (06.10 — первый день без процедуры).
        stepper_row = first_row + CARD_LABELS.index("Степпер")
        fills = [sheet[f"{column(i)}{stepper_row}"].fill.fill_type for i in range(6)]
        assert fills == ["solid", "solid", None, None, "solid", "solid"]
        assert all(sheet[f"{column(i)}{stepper_row}"].value is None for i in range(first_block))
        # Верхний блок «Расписание занятий» — типичный день: время настоящим значением Excel.
        top = [sheet[f"B{row}"].value for row in range(4, 14) if sheet[f"B{row}"].value is not None]
        assert top, "верхний блок заполнен расписанием (фаза 2)"
        assert all(isinstance(value, dt_time) for value in top)
        assert top == sorted(top)
        assert CardExport.objects.get().program == program

    def test_shrm5_second_block(self, client, doctor):
        client.force_login(doctor)
        client.post(IMPORT_URL, {"file": upload(shrm=5, rows=FULL_ROWS)})
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]))

        sheet = load_workbook(BytesIO(response.content)).active
        second = [sheet[f"{column(i)}31"].value for i in range(10)]
        assert all(isinstance(value, datetime) for value in second)
        assert [value.date() for value in second] == program.course_dates()[10:]
        assert [sheet[f"A{32 + i}"].value for i in range(len(CARD_LABELS))] == CARD_LABELS

    def test_unrecognized_after_import_blocks_card(self, client, doctor):
        client.force_login(doctor)
        client.post(
            IMPORT_URL,
            {"file": upload(rows=[("02.10.26", "Неизвестная процедура 10 мин", "")])},
        )
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]), follow=True)

        assert response.redirect_chain[-1][0] == reverse("programs:detail", args=[program.pk])
        assert "Неизвестная процедура 10 мин" in response.content.decode()

    def test_too_many_rows_for_shrm3_is_a_message(self, client, doctor):
        client.force_login(doctor)
        rows = [("02.10.26", "Имитрон", "")] * 16
        client.post(IMPORT_URL, {"file": upload(shrm=3, rows=rows)})
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]), follow=True)

        assert "Процедур для карты 16" in response.content.decode()
        assert not CardExport.objects.exists()

    @pytest.mark.parametrize(
        ("fio", "surname"),
        [
            ("О'Тестова Анна Сергеевна", "О'Тестова"),
            ("Тестова-Пробная Анна Сергеевна", "Тестова-Пробная"),
        ],
    )
    def test_filename_from_sheet(self, client, doctor, fio, surname):
        client.force_login(doctor)
        client.post(IMPORT_URL, {"file": upload(fio=fio)})
        program = Program.objects.get()

        response = client.get(reverse("cards:download", args=[program.pk]))

        disposition = response["Content-Disposition"]
        encoded = disposition.split("filename*=UTF-8''", 1)[1]
        assert "'" not in encoded and " " not in encoded
        assert unquote(encoded) == f"Программа_{surname}_5402_02.10.2026.xlsx"


class TestExtendedCourseCard:
    """Продлённый курс печатается на наименьшем шаблоне, где хватает дат (TZ.md, решение 40)."""

    @pytest.mark.parametrize(
        ("shrm", "days", "expected"),
        [(3, 11, 3), (3, 13, 4), (4, 15, 4), (4, 18, 5), (3, 20, 5), (5, 20, 5), (5, 10, 5)],
    )
    def test_choose_template(self, shrm, days, expected):
        assert choose_template(shrm, days) == expected

    def test_longer_than_any_template(self):
        with pytest.raises(CardRenderError, match="Дней курса 25"):
            choose_template(4, 25)

    def test_shrm4_extended_prints_on_shrm5_template(self, doctor, make_program, procedures):
        program = make_program(shrm=4)  # 05.10–19.10
        program.end_date = date(2026, 10, 22)  # продлили до 18 дней

        save_program(doctor, program, end_date_changed=True)
        Prescription.objects.create(program=program, procedure=procedures["group"])

        _, content = export_card(doctor, program)
        sheet = load_workbook(BytesIO(content)).active

        assert sheet.title == "ШРМ 5"
        assert sheet["C18"].value.date() == date(2026, 10, 5)
        assert sheet["J31"].value.date() == date(2026, 10, 22), "18-я дата во втором блоке"
        assert sheet["K31"].value is None
