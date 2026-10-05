"""Перестройка при изменении смены и распорядка (TZ.md FR-SCH-13, FR-STF-6, решение 43),
отчёт «Загрузка инструкторов» (FR-SCH-17), публичная шахматка (FR-ACC-5), выгрузка (FR-CRD-7).

Справочник — полный, инструкторы — вымышленные из seed_dev.
"""

from datetime import date, time, timedelta
from io import BytesIO, StringIO

import pytest
from django.core.management import call_command
from django.urls import reverse
from openpyxl import load_workbook

from apps.catalog.models import InstructorSlot, Procedure
from apps.programs import services as programs
from apps.programs.models import Prescription, Program
from apps.scheduling import board, manual, reports, staff_changes
from apps.scheduling.models import Booking, StaffChange
from apps.scheduling.services import choose_preferred_instructor
from apps.staff import services as staff_services
from apps.staff.models import Instructor, InstructorBlock, InstructorDuty, ShiftException

pytestmark = pytest.mark.usefixtures("full_catalog")

MONDAY = date(2026, 10, 5)
INDIVIDUAL = "Индивидуальное занятие"


@pytest.fixture
def staff(settings, doctor, rehab, admin_user) -> dict[str, Instructor]:
    settings.DEBUG = True
    call_command("seed_dev", stdout=StringIO())
    return {item.short_name: item for item in Instructor.objects.all()}


def prescribe(user, program: Program, name: str, **fields) -> Prescription:
    procedure = Procedure.objects.get(name=name, department=None)
    return programs.add_prescription(
        user, Prescription(program=program, procedure=procedure, **fields)
    )


@pytest.fixture
def patient(doctor, staff, make_program) -> Program:
    program = make_program(full_name="Иванов Иван Иванович", room="9")
    prescribe(doctor, program, INDIVIDUAL)
    return program


def individual(program: Program, day: date = MONDAY) -> Booking:
    return program.bookings.get(kind="INDIVIDUAL", date=day)


def sick(user, instructor: Instructor, day: date = MONDAY) -> StaffChange | None:
    return staff_changes.set_not_working(user, instructor, day)


class TestRebuild:
    def test_sick_instructor_replaced_on_that_date_only(self, rehab, patient):
        before = individual(patient)
        tuesday = individual(patient, MONDAY + timedelta(days=1))

        change = sick(rehab, before.instructor)

        after = individual(patient)
        assert after.instructor != before.instructor
        assert after.start == before.start, "замена — в тот же слот к другому инструктору"
        assert individual(patient, MONDAY + timedelta(days=1)).instructor == tuesday.instructor
        assert len(change.items) == 1
        item = change.items[0]
        assert item["label"] == "9п Иванов" and item["date"] == "2026-10-05"
        assert item["before"]["instructor"] == before.instructor.short_name
        assert item["after"]["instructor"] == after.instructor.short_name
        assert not item["pinned"] and not item["preferred"]
        patient.refresh_from_db()
        assert "DAY_DEVIATION" in [i["code"] for i in patient.schedule_issues]

    def test_pinned_is_moved_and_stays_pinned(self, rehab, patient):
        booking = individual(patient)
        target = next(
            (i, s)
            for i, s in board.free_seats(MONDAY)
            if not s.is_evening and s.start >= time(13) and i != booking.instructor
        )
        board.place_patient(rehab, patient, *target, MONDAY)

        change = sick(rehab, target[0])

        moved = individual(patient)
        assert moved.instructor != target[0] and moved.pinned
        assert change.items[0]["pinned"]

    def test_preferred_is_marked(self, rehab, patient):
        instructor = individual(patient).instructor
        choose_preferred_instructor(rehab, patient, instructor)

        change = sick(rehab, instructor)

        assert change.items[0]["preferred"]

    def test_nothing_affected_gives_no_summary(self, rehab, patient, staff):
        free = next(
            name for name in ("Зайцев", "Кузнецова", "Морозова")
            if staff[name] != individual(patient).instructor
        )  # fmt: skip

        assert sick(rehab, staff[free]) is None
        assert ShiftException.objects.filter(instructor=staff[free], date=MONDAY).exists()

    def test_unplaced_when_nobody_is_free(self, rehab, patient, staff):
        own = individual(patient).instructor
        for item in Instructor.objects.exclude(pk=own.pk):
            ShiftException.objects.create(instructor=item, date=MONDAY, is_working=False)

        change = sick(rehab, own)

        assert change.items[0]["after"] is None and change.unplaced == 1
        assert not patient.bookings.filter(kind="INDIVIDUAL", date=MONDAY).exists()
        patient.refresh_from_db()
        assert any(i["code"] == "INDIVIDUAL_UNPLACED" for i in patient.schedule_issues)

    def test_block_and_duty_trigger_rebuild(self, rehab, patient):
        booking = individual(patient)
        block = InstructorBlock(
            instructor=booking.instructor, date=MONDAY, slot=booking.slot, kind="BOS"
        )
        staff_services.save_block(rehab, block)

        change = staff_changes.after_block(rehab, block)

        assert change is not None and individual(patient).instructor != booking.instructor
        tuesday = individual(patient, MONDAY + timedelta(days=1))
        duty = staff_services.save_duty(
            rehab,
            InstructorDuty(
                instructor=tuesday.instructor,
                slot=tuesday.slot,
                kind="METHOD_WORK",
                valid_from=MONDAY + timedelta(days=1),
            ),
        )
        change = staff_changes.after_duty(rehab, duty)
        assert change is not None
        assert {item["date"] for item in change.items} >= {"2026-10-06"}
        later = patient.bookings.filter(
            kind="INDIVIDUAL", date__gt=MONDAY, instructor__isnull=False
        )
        assert all(
            not (b.instructor_id == tuesday.instructor_id and b.slot_id == tuesday.slot_id)
            for b in later
        )


class TestRevert:
    def test_revert_after_exception_removed(self, rehab, patient):
        before = individual(patient)
        change = sick(rehab, before.instructor)
        staff_services.set_working(rehab, before.instructor, MONDAY, True)

        problems = staff_changes.revert(rehab, change)

        assert problems == []
        back = individual(patient)
        assert (back.instructor, back.slot) == (before.instructor, before.slot)
        change.refresh_from_db()
        assert change.reverted_at is not None and change.reverted_by == rehab
        with pytest.raises(manual.EditError, match="Уже возвращено"):
            staff_changes.revert(rehab, change)

    def test_revert_while_still_sick_reports_problem(self, rehab, patient):
        before = individual(patient)
        change = sick(rehab, before.instructor)

        problems = staff_changes.revert(rehab, change)

        assert len(problems) == 1 and "9п Иванов, 05.10" in problems[0]
        assert individual(patient).instructor != before.instructor

    def test_revert_unplaced_recreates_booking(self, rehab, patient, staff):
        own = individual(patient).instructor
        others = list(Instructor.objects.exclude(pk=own.pk))
        for item in others:
            ShiftException.objects.create(instructor=item, date=MONDAY, is_working=False)
        change = sick(rehab, own)
        staff_services.set_working(rehab, own, MONDAY, True)

        assert staff_changes.revert(rehab, change) == []
        assert individual(patient).instructor == own

    def test_doctor_cannot_revert(self, doctor, rehab, patient):
        change = sick(rehab, individual(patient).instructor)

        with pytest.raises(manual.EditError):
            staff_changes.revert(doctor, change)


class TestViews:
    def test_shift_toggle_shows_notice(self, client, rehab, patient):
        client.force_login(rehab)
        instructor = individual(patient).instructor

        html = client.post(
            reverse("staff:shift_toggle", args=[instructor.pk]),
            {"day": "2026-10-05"},
            headers={"HX-Request": "true"},
        ).content.decode()

        change = StaffChange.objects.get()
        assert "Перестроено занятий: 1" in html
        assert reverse("scheduling:change", args=[change.pk]) in html

    def test_board_off_redirects_to_summary(self, client, rehab, patient):
        client.force_login(rehab)
        instructor = individual(patient).instructor

        response = client.post(
            reverse("scheduling:board_off"), {"instructor": instructor.pk, "date": "2026-10-05"}
        )

        change = StaffChange.objects.get()
        assert response.url == reverse("scheduling:change", args=[change.pk])
        html = client.get(response.url).content.decode()
        assert "9п Иванов" in html and "Вернуть как было" in html

    def test_board_off_without_patients_returns_to_board(self, client, rehab, patient, staff):
        client.force_login(rehab)
        free = next(
            name for name in ("Зайцев", "Кузнецова")
            if staff[name] != individual(patient).instructor
        )  # fmt: skip

        response = client.post(
            reverse("scheduling:board_off"), {"instructor": staff[free].pk, "date": "2026-10-05"}
        )

        assert response.url.endswith("?date=2026-10-05")

    def test_revert_view_and_list(self, client, rehab, patient):
        instructor = individual(patient).instructor
        change = sick(rehab, instructor)
        staff_services.set_working(rehab, instructor, MONDAY, True)
        client.force_login(rehab)

        assert change.title in client.get(reverse("scheduling:changes")).content.decode()
        html = client.post(
            reverse("scheduling:change", args=[change.pk]), follow=True
        ).content.decode()
        assert "Занятия возвращены на прежние места." in html and "Возвращено" in html

    def test_doctor_cannot_open_summary_or_board_off(self, client, doctor, rehab, patient):
        change = sick(rehab, individual(patient).instructor)
        client.force_login(doctor)

        assert client.get(reverse("scheduling:change", args=[change.pk])).status_code == 403
        response = client.post(
            reverse("scheduling:board_off"), {"instructor": 1, "date": "2026-10-05"}
        )
        assert response.status_code == 403


# --- Отчёт «Загрузка инструкторов» ----------------------------------------------------------


class TestLoad:
    def test_counts_weekday_sessions_and_working_days(self, rehab, patient, staff):
        owner = individual(patient).instructor

        report = reports.instructor_load(rehab, MONDAY, MONDAY + timedelta(days=6))

        rows = {row.name: row for row in report.rows if not row.is_total}
        assert rows[owner.short_name].sessions == 5, "5 будних дней, выходные не в счёт"
        assert rows["Зайцев"].days == 5 and rows["Зайцев"].sessions == 0
        pair = next(row for row in report.rows if row.is_total and row.name == "Волков/ Лебедева")
        assert pair.days == rows["Волков"].days + rows["Лебедева"].days
        assert rows[owner.short_name].average == 1

    def test_reversed_period_and_rights(self, doctor, rehab, staff):
        report = reports.instructor_load(rehab, MONDAY, MONDAY - timedelta(days=3))
        assert report.start < report.end
        with pytest.raises(Exception, match="специалист ФР"):
            reports.instructor_load(doctor, MONDAY, MONDAY)

    def test_view(self, client, rehab, patient):
        client.force_login(rehab)

        html = client.get(
            reverse("scheduling:load"), {"start": "2026-10-05", "end": "2026-10-09"}
        ).content.decode()

        assert "Загрузка инструкторов" in html and "Волков/ Лебедева — вместе" in html


# --- Публичная шахматка и выгрузка ----------------------------------------------------------


class TestPublicBoard:
    def test_open_without_login_shows_department_and_no_links(self, client, patient):
        response = client.get(reverse("public_board"), {"date": "2026-10-05"})

        html = response.content.decode()
        assert response.status_code == 200
        assert "9п ОМР № 4 Иванов" in html
        assert "Иван Иванович" not in html
        assert reverse("programs:detail", args=[patient.pk]) not in html
        assert "board/cell" not in html and "board/export" not in html

    def test_networks_limit_access(self, client, settings, patient):
        settings.PUBLIC_BOARD_NETWORKS = ["10.0.0.0/8"]

        assert client.get(reverse("public_board")).status_code == 404
        assert client.get(reverse("public_board"), REMOTE_ADDR="10.1.2.3").status_code == 200

    def test_bad_network_setting_denies(self, client, settings, staff):
        settings.PUBLIC_BOARD_NETWORKS = ["не сеть"]

        assert client.get(reverse("public_board")).status_code == 404


class TestExport:
    def test_xlsx_values_types_and_layout(self, client, doctor, rehab, patient, staff):
        prescribe(doctor, patient, "st-150")
        client.force_login(rehab)

        response = client.get(reverse("scheduling:board_export"), {"date": "2026-10-05"})

        assert response["Content-Type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml"
        )
        assert "filename*=UTF-8''" in response["Content-Disposition"]
        sheet = load_workbook(BytesIO(response.content)).active
        assert sheet["A1"].is_date and sheet["A1"].value.date() == MONDAY
        assert sheet["A1"].number_format == "dd.mm.yyyy"
        headers = [c.value for c in sheet[1]]
        assert "Волков/ Лебедева" in headers and "Соколов" in headers
        slots = list(InstructorSlot.objects.order_by("start"))
        assert sheet["A2"].value == f"{slots[0].start.hour}:{slots[0].start:%M}-" + (
            f"{slots[0].end.hour}:{slots[0].end:%M}"
        )
        booking = individual(patient)
        column = headers.index(booking.instructor.team_label) + 1
        row = slots.index(booking.slot) + 2
        assert sheet.cell(row, column).value == "9п Иванов"
        method = headers.index("Соколов") + 1
        busy = sheet.cell(slots.index(InstructorSlot.objects.get(start=time(15))) + 2, method)
        assert busy.value == "Метод. работа" and busy.fill.start_color.rgb == "FFD9D9D9"
        # Тренажёры — через пустую колонку, время — настоящее время Excel.
        time_col = headers.index("Время") + 1
        assert headers[time_col - 2] is None
        assert sheet.cell(2, time_col).number_format == "h:mm"
        assert isinstance(sheet.cell(2, time_col).value, time)
        st = headers.index("st-150") + 1
        values = [sheet.cell(r, st).value for r in range(2, sheet.max_row + 1)]
        assert "9п Иванов" in values
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1 and sheet.sheet_properties.pageSetUpPr.fitToPage
