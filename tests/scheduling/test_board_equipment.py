"""Тренажёры в шахматке: перенос перетаскиванием на другое время (TZ.md FR-SCH-10a).

Сегодня в тестах — вторник 06.10.2026; окно st-150 — 13:00–15:00, шаг 15 минут, вместимость 1.
"""

import json
from datetime import date, time, timedelta

import pytest
from django.urls import reverse

from apps.accounts.models import Preset
from apps.programs.models import Prescription
from apps.programs.services import add_prescription
from apps.scheduling import board, board_days
from apps.scheduling.models import Booking, BookingKind
from apps.scheduling.services import replan

MON, TUE, WED = date(2026, 10, 5), date(2026, 10, 6), date(2026, 10, 7)


@pytest.fixture(autouse=True)
def today(monkeypatch):
    monkeypatch.setattr(board_days, "today", lambda: TUE)


@pytest.fixture
def patient(doctor, make_program, procedures):
    """make: пациент со st-150 (курс с 05.10)."""

    def make(room: str = "9", **fields):
        program = make_program(room=room, full_name=f"Пациентов{room} Тест Тестович", **fields)
        add_prescription(doctor, Prescription(program=program, procedure=procedures["equipment"]))
        return program

    return make


def equipment_on(program, day: date) -> Booking:
    return Booking.objects.get(program=program, date=day, kind=BookingKind.EQUIPMENT)


@pytest.fixture
def two(patient):
    board_days.ensure_boards()
    first, second = patient("1"), patient("2")
    assert equipment_on(first, TUE).start != equipment_on(second, TUE).start
    return first, second


class TestService:
    def test_move_to_free_time(self, two, rehab):
        first, _second = two
        booking = equipment_on(first, TUE)

        board.move_equipment(rehab, booking, booking.version, time(14, 30))

        moved = equipment_on(first, TUE)
        assert (moved.start, moved.end, moved.pinned) == (time(14, 30), time(14, 45), True)
        # Только эта дата: в среду время прежнее.
        assert equipment_on(first, date(2026, 10, 7)).start == booking.start

    def test_over_capacity_needs_reason_and_keeps_it_as_note(self, two, rehab):
        first, second = two
        target = equipment_on(first, TUE).start
        booking = equipment_on(second, TUE)

        with pytest.raises(board.ConfirmationRequired):
            board.move_equipment(rehab, booking, booking.version, target)
        board.move_equipment(rehab, booking, booking.version, target, "держит 15 мин")

        moved = equipment_on(second, TUE)
        assert (moved.start, moved.note) == (target, "держит 15 мин")
        row = next(r for r in board.board(rehab, TUE).equipment_rows if r.start == target)
        _item, patients, over = row.cells[0]
        assert [p.label for p in patients] == ["1п Пациентов1", "2п Пациентов2 держит 15 мин"]
        assert over

    def test_outside_window_needs_reason(self, two, rehab):
        booking = equipment_on(two[0], TUE)
        with pytest.raises(board.ConfirmationRequired, match="вне окна"):
            board.move_equipment(rehab, booking, booking.version, time(16, 0))

    def test_only_today_and_tomorrow(self, patient, rehab):
        board_days.ensure_boards()
        booking = equipment_on(patient(start_date=date(2026, 10, 2)), MON)
        with pytest.raises(board.BoardError, match="сегодняшнее и завтрашнее"):
            board.move_equipment(rehab, booking, booking.version, time(14, 0))

    def test_rights(self, two, doctor, make_user, other_department):
        booking = equipment_on(two[0], TUE)
        with pytest.raises(board.BoardError, match="Нет права"):
            board.move_equipment(doctor, booking, booking.version, time(14, 0))
        stranger = make_user("rehab2", (other_department, Preset.REHAB))
        with pytest.raises(board.BoardError, match="другого отделения"):
            board.move_equipment(stranger, booking, booking.version, time(14, 0))

    def test_stale_version(self, two, rehab):
        booking = equipment_on(two[0], TUE)
        with pytest.raises(board.BoardError, match="изменена"):
            board.move_equipment(rehab, booking, booking.version + 1, time(14, 0))

    def test_sheet_keeps_labels(self, two, rehab):
        result = board.sheet(board.board(rehab, TUE))
        labels = {name for row in result.equipment_cells for names, _over in row for name in names}
        assert labels == {"1п Пациентов1", "2п Пациентов2"}


class TestScreen:
    def url(self, booking: Booking) -> str:
        return reverse("scheduling:board_equipment_move", args=[booking.pk])

    def test_board_marks_draggable_equipment(self, client, two, rehab):
        client.force_login(rehab)
        page = client.get(reverse("scheduling:board"), {"date": "2026-10-06"}).content.decode()
        booking = equipment_on(two[0], TUE)
        assert f'data-booking="{booking.pk}"' in page and "data-equipment-drop" in page
        # Многострочный комментарий шаблона не попадает на страницу.
        assert "Шахматка на дату (FR-SCH-14)" not in page

    def test_doctor_sees_no_drag(self, client, two, doctor):
        client.force_login(doctor)
        page = client.get(reverse("scheduling:board"), {"date": "2026-10-06"}).content.decode()
        assert "data-equipment-drop" not in page
        booking = equipment_on(two[0], TUE)
        response = client.post(self.url(booking), {"start": "14:00", "version": booking.version})
        assert response.status_code == 403

    def test_drop_moves_and_returns_board(self, client, two, rehab):
        client.force_login(rehab)
        booking = equipment_on(two[0], TUE)

        page = client.post(
            self.url(booking), {"start": "14:30", "version": booking.version}
        ).content.decode()

        assert "Время тренажёра изменено" in page and 'hx-swap-oob="true"' in page
        assert equipment_on(two[0], TUE).start == time(14, 30)

    def test_over_capacity_asks_reason(self, client, two, rehab):
        client.force_login(rehab)
        first, second = two
        booking = equipment_on(second, TUE)
        start = f"{equipment_on(first, TUE).start:%H:%M}"

        page = client.post(self.url(booking), {"start": start, "version": booking.version})

        assert page.status_code == 200
        text = page.content.decode()
        assert "при вместимости 1" in text and 'name="reason"' in text
        assert equipment_on(second, TUE).start != equipment_on(first, TUE).start
        client.post(
            self.url(booking), {"start": start, "version": booking.version, "reason": "по просьбе"}
        )
        assert equipment_on(second, TUE).note == "по просьбе"

    def test_bad_time_is_404(self, client, two, rehab):
        client.force_login(rehab)
        booking = equipment_on(two[0], TUE)
        response = client.post(self.url(booking), {"start": "abc", "version": booking.version})
        assert response.status_code == 404


def test_board_script_has_version(client, rehab):
    """Скрипт шахматки подключается с версией — после правки браузер не возьмёт старый."""
    client.force_login(rehab)
    page = client.get(reverse("scheduling:board"), {"date": "2026-10-06"}).content.decode()
    assert "js/board.js?v=" in page and "css/app.css?v=" in page


class TestBusyHighlight:
    """Подсветка времени, где у пациента другие занятия (FR-SCH-10a)."""

    def test_patient_carries_other_lessons_of_the_day(self, two, rehab):
        first, _second = two
        own = equipment_on(first, TUE)
        result = board.board(rehab, TUE)
        chip = next(
            p for row in result.equipment_rows for _i, ps, _o in row.cells for p in ps
            if p.booking.pk == own.pk
        )  # fmt: skip
        # Переносимое занятие не мешает самому себе.
        assert json.loads(chip.busy) == []
        other = board.occupied({first.pk}, TUE)[first.pk]
        assert json.loads(board.busy_json(other)) == [
            [13 * 60 + own.start.minute, 13 * 60 + own.start.minute + 15,
             f"{own.start:%-H:%M} st-150"],
        ]  # fmt: skip

    def test_screen_has_times_and_busy(self, client, two, rehab):
        client.force_login(rehab)
        page = client.get(reverse("scheduling:board"), {"date": "2026-10-06"}).content.decode()
        assert 'data-from="13:00" data-minutes="15"' in page and "data-busy=" in page

    def test_public_board_has_no_busy(self, client, two, settings):
        settings.PUBLIC_BOARD_NETWORKS = []
        response = client.get(reverse("public_board"), {"date": "2026-10-06"})
        assert response.status_code == 200
        page = response.content.decode()
        assert "1п" in page and "data-busy=" not in page


class TestNote:
    """Пометка у записи на тренажёр — как у индивидуальных (FR-SCH-10)."""

    def test_note_goes_to_following_days_and_survives_replan(self, two, rehab):
        first, _second = two
        booking = equipment_on(first, TUE)

        board.set_equipment_note(rehab, booking, booking.version, "  с  утяжелением ")

        saved = equipment_on(first, TUE)
        assert (saved.note, saved.pinned, saved.start) == ("с утяжелением", False, booking.start)
        notes = dict(first.bookings.filter(kind=BookingKind.EQUIPMENT).values_list("date", "note"))
        # С этой даты и дальше, как у индивидуальных; прошлое не меняется.
        assert all(note == "" for day, note in notes.items() if day < TUE)
        assert {day: note for day, note in notes.items() if day >= TUE} == {
            day: "с утяжелением" for day in notes if day >= TUE
        }
        replan(first)
        assert equipment_on(first, WED).note == "с утяжелением"
        labels = {p.label for row in board.board(rehab, TUE).equipment_rows
                  for _i, ps, _o in row.cells for p in ps}  # fmt: skip
        assert "1п Пациентов1 с утяжелением" in labels

    def test_tomorrow_note_leaves_today(self, two, rehab):
        booking = equipment_on(two[0], WED)
        board.set_equipment_note(rehab, booking, booking.version, "art")
        assert equipment_on(two[0], TUE).note == ""
        assert equipment_on(two[0], date(2026, 10, 8)).note == "art"

    def test_new_course_day_inherits(self, two, rehab):
        first = two[0]
        last = first.end_date
        booking = equipment_on(first, TUE)
        board.set_equipment_note(rehab, booking, booking.version, "art")
        first.end_date = last + timedelta(days=2)
        first.save(update_fields=["end_date"])

        replan(first)

        added = first.bookings.filter(kind=BookingKind.EQUIPMENT, date__gt=last - timedelta(1))
        assert added.exists() and set(added.values_list("note", flat=True)) == {"art"}

    def test_capacity_reason_stays_on_its_day(self, two, rehab):
        first, second = two
        booking = equipment_on(second, TUE)
        board.move_equipment(
            rehab, booking, booking.version, equipment_on(first, TUE).start, "держит 15 мин"
        )
        replan(second)
        assert equipment_on(second, TUE).note == "держит 15 мин"
        assert equipment_on(second, WED).note == ""

    def test_empty_clears(self, two, rehab):
        booking = equipment_on(two[0], TUE)
        board.set_equipment_note(rehab, booking, booking.version, "art")
        booking = equipment_on(two[0], TUE)
        board.set_equipment_note(rehab, booking, booking.version, " ")
        assert equipment_on(two[0], TUE).note == ""
        assert equipment_on(two[0], WED).note == ""

    def test_rights_version_and_date(
        self, two, patient, doctor, rehab, make_user, other_department
    ):
        booking = equipment_on(two[0], TUE)
        with pytest.raises(board.BoardError, match="Нет права"):
            board.set_equipment_note(doctor, booking, booking.version, "x")
        stranger = make_user("rehab2", (other_department, Preset.REHAB))
        with pytest.raises(board.BoardError, match="другого отделения"):
            board.set_equipment_note(stranger, booking, booking.version, "x")
        with pytest.raises(board.BoardError, match="изменена"):
            board.set_equipment_note(rehab, booking, booking.version + 1, "x")
        past = equipment_on(patient("5", start_date=date(2026, 10, 2)), MON)
        with pytest.raises(board.BoardError, match="сегодняшнее и завтрашнее"):
            board.set_equipment_note(rehab, past, past.version, "x")

    def test_screen(self, client, two, rehab, doctor):
        booking = equipment_on(two[0], TUE)
        client.force_login(rehab)
        page = client.get(reverse("scheduling:board"), {"date": "2026-10-06"}).content.decode()
        assert reverse("scheduling:board_equipment", args=[booking.pk]) in page

        panel = client.get(reverse("scheduling:board_equipment", args=[booking.pk]))
        assert panel.status_code == 200 and 'name="note"' in panel.content.decode()

        url = reverse("scheduling:board_equipment_note", args=[booking.pk])
        done = client.post(url, {"note": "(шар)", "version": booking.version}).content.decode()
        assert "Пометка сохранена" in done and 'hx-swap-oob="true"' in done
        assert "Пациентов1 (шар)" in done

        stale = client.post(url, {"note": "x", "version": booking.version}).content.decode()
        assert "изменена" in stale and 'name="note"' in stale

        client.force_login(doctor)
        assert (
            client.get(reverse("scheduling:board_equipment", args=[booking.pk])).status_code == 403
        )
        assert client.post(url, {"note": "x", "version": 0}).status_code == 403


class TestNoteInCard:
    """Пометка тренажёра печатается в блоке «Расписание занятий» карты (FR-CRD-2)."""

    def test_card_has_note_of_nearest_day(self, two, rehab, doctor):
        from io import BytesIO

        from openpyxl import load_workbook

        from apps.cards.services import build_card_data, export_card

        first = two[0]
        booking = equipment_on(first, WED)
        board.set_equipment_note(rehab, booking, booking.version, "(шар)")
        # Сегодня (вторник) пометки ещё нет — в карту идёт ближайший день, т. е. сегодняшний.
        assert [i.label for i in build_card_data(first).schedule if "st-150" in i.label] == [
            "st-150"
        ]
        today = equipment_on(first, TUE)
        board.set_equipment_note(rehab, today, today.version, "(шар)")

        assert [i.label for i in build_card_data(first).schedule if "st-150" in i.label] == [
            "st-150 (шар)"
        ]
        _name, content = export_card(doctor, first)
        values = [
            cell.value for row in load_workbook(BytesIO(content)).active.iter_rows() for cell in row
        ]
        assert "st-150 (шар)" in values
